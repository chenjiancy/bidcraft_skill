# -*- coding: utf-8 -*-
"""⑨ 最小重现：docxtpl 0.20.2 行循环（{%tr %}）是否展开。"""
import sys
from pathlib import Path

OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(OUT / "_deps"))

from docx import Document
from docx.shared import Pt
from docxtpl import DocxTemplate

tpl = OUT / "mini_tpl.docx"
doc = Document()
tbl = doc.add_table(rows=3, cols=2)
c00 = tbl.rows[0].cells[0].paragraphs[0]
c00.add_run("{%tr for row in rows %}")
tbl.rows[0].cells[1].paragraphs[0].add_run("{{ row.姓名 }}")
tbl.rows[1].cells[0].paragraphs[0].add_run("{{ row.备注 }}")
tbl.rows[2].cells[0].paragraphs[0].add_run("{%tr endfor %}")
doc.save(str(tpl))

t = DocxTemplate(str(tpl))
t.render({
    "rows": [{"姓名": "A", "备注": "x"}, {"姓名": "B", "备注": "y"}],
})
t.save(str(OUT / "mini_out.docx"))

d = Document(str(OUT / "mini_out.docx"))
print("rows:", len(d.tables[0].rows))
for r in d.tables[0].rows:
    print([c.text for c in r.cells])
