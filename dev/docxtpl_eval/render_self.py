# -*- coding: utf-8 -*-
"""⑨ docxtpl 对比评估 · 路线 A：自研渲染（docx_util + 行 deepcopy，即现有产品路径的公共工具）。"""
import copy
import sys
from pathlib import Path

OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(OUT.parent.parent / "scripts"))   # E:\bidcraft_skill\scripts，使 _shared 可导入

from docx import Document  # noqa: E402

from _shared import docx_util  # noqa: E402

DATA = {
    "project_name": "示例示例园区尾水水质提升工程（EPC总承包）监理",
    "rows": [
        {"姓名": "张三", "出生年月": "1970年5月", "备注": "总监"},
        {"姓名": "赵六", "出生年月": "1982年11月", "备注": "专监"},
        {"姓名": "赵七", "出生年月": "1990年3月", "备注": "监理员"},
    ],
}


def render_self(tpl_path, out_path):
    doc = Document(tpl_path)
    # 1) 段落占位（跨 run 安全：docx_util.set_para_text 合并 run 保留首 run 样式）
    for p in doc.paragraphs:
        text = docx_util.para_text(p)
        if "【项目名称】" in text:
            docx_util.set_para_text(p, text.replace("【项目名称】", DATA["project_name"]))
    # 2) 表格范本行复制为 N 行 + 逐格替换
    tbl = doc.tables[0]
    tmpl_row = tbl.rows[1]
    n = len(DATA["rows"])
    if n > 1:
        node = tmpl_row._tr
        last = node
        for _ in range(n - 1):
            new = copy.deepcopy(node)
            last.addnext(new)
            last = new
    rows = tbl.rows[1:1 + n]
    for ri, rec in enumerate(DATA["rows"]):
        for ci, key in enumerate(("姓名", "出生年月", "备注")):
            cell = rows[ri].cells[ci]
            t = docx_util.para_text(cell.paragraphs[0])
            docx_util.set_para_text(cell.paragraphs[0], t.replace("【%s】" % key, rec[key]))
    doc.save(str(out_path))
    return str(out_path)


if __name__ == "__main__":
    print(render_self(OUT / "tpl_self.docx", OUT / "out_self.docx"))
