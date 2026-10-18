# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, r'E:\bidcraft_skill\dev')
from docx import Document

base = r'E:\天翼同步盘\投标文件\2026年\和县2026年老旧小区改造项目（EPC总承包）监理采购\和县'
words = []
for dp, dn, fn in os.walk(base):
    for f in fn:
        if f.endswith('.docx') and not f.startswith('~'):
            words.append(os.path.join(dp, f))

for f in words:
    fn = os.path.basename(f)
    try:
        doc = Document(f)
        breaks = []
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip() == '':
                if 'w:br' in p._p.xml and 'w:type="page"' in p._p.xml:
                    breaks.append(i)
        sects = len(doc.sections)
        first_txt = ''
        for p in doc.paragraphs:
            if p.text.strip():
                first_txt = p.text.strip()[:22]; break
        print('%-34s 首段[%s] 节数=%d 分页符=%d 总段=%d' % (fn, first_txt, sects, len(breaks), len(doc.paragraphs)))
    except Exception as e:
        print('%-34s 失败 %s' % (fn, e))
