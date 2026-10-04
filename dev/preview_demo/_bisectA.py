# -*- coding: utf-8 -*-
"""二分A：预览框=空段落（无图）→ 重新生成资格证明 → Word 测试。
二分B：若 A 通过，则问题在图片绘制 XML；提取一张预览图 XML 供人工检查。"""
import shutil
import sys
from pathlib import Path

sys.path.insert(0, r"E:\bidcraft_skill\scripts")

DEMO = Path(r"E:\bidcraft_skill\dev\preview_demo")
OUT = DEMO / "triage" / "gen_nopic_资格证明.docx"

from m5_project import generator as gen  # noqa: E402

# 打补丁：_make_ph_preview_para 返回空段落（无图、无 marker）
orig = gen._make_ph_preview_para
def nopic(doc, ph_text):
    p = doc.add_paragraph()
    p.alignment = 1
    return p
gen._make_ph_preview_para = nopic

ENT = Path(r"E:\监理标书制作\和县建设工程监理有限公司")
PROJ = "马鞍山和县化工园尾水水质提升工程（EPC总承包）监理"
TMPDIR = DEMO / "triage" / "gen_nopic"
if TMPDIR.exists():
    shutil.rmtree(TMPDIR)
TMPDIR.mkdir(parents=True)

try:
    res = gen.generate(ENT, PROJ, out_dir=TMPDIR, register_baseline=False)
    src = TMPDIR / "资格证明及辅助资料表.docx"
    shutil.copy(src, OUT)
    print("generated", res["文件"][0]["图片占位框"], "预览框(应为0)")
finally:
    gen._make_ph_preview_para = orig

# Word 导出测试
import win32com.client  # noqa: E402
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
word.DisplayAlerts = 0
try:
    doc = word.Documents.Open(str(OUT), ReadOnly=True)
    doc.ExportAsFixedFormat(str(OUT.with_suffix(".pdf")), 17)
    doc.Close(False)
    print("Word OK: no-picture previews 版本")
except Exception as e:
    print("Word FAIL: no-picture previews 版本 ->", (getattr(e, "excepinfo", None) or [""])[2])
finally:
    word.Quit()

# 提取一张预览图 XML
from docx import Document  # noqa: E402
from lxml import etree  # noqa: E402
full = Document(str(DEMO / "项目模板" / "资格证明及辅助资料表.docx"))
for p in full.paragraphs:
    descrs = [el.get("descr") or "" for el in p._p.iter() if el.tag.endswith("}docPr")]
    if descrs:
        xml = etree.tostring(p._p, pretty_print=True).decode("utf-8")
        (DEMO / "triage" / "preview_para_sample.xml").write_text(xml, encoding="utf-8")
        print("preview xml sample saved, len", len(xml))
        break
