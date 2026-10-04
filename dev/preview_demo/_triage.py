# -*- coding: utf-8 -*-
"""定位：旧模板 vs 新模板往返保存 vs 去掉预览框。"""
import shutil
import sys
from pathlib import Path

sys.path.insert(0, r"E:\bidcraft_skill\scripts")
from docx import Document  # noqa: E402

DEMO = Path(r"E:\bidcraft_skill\dev\preview_demo")
TMP = DEMO / "triage"
TMP.mkdir(exist_ok=True)

OLD = Path(r"E:\监理标书制作\示例建设工程监理有限公司\项目级\示例示例园区尾水水质提升工程（EPC总承包）监理\项目模板\资格证明及辅助资料表.docx")
NEW = DEMO / "项目模板" / "资格证明及辅助资料表.docx"

# 1) 旧模板复制到 triage
shutil.copy(OLD, TMP / "old_资格证明.docx")

# 2) 新模板往返保存（python-docx 打开再保存）
d = Document(str(NEW))
rt = TMP / "new_roundtrip_资格证明.docx"
d.save(str(rt))

# 3) 新模板去掉所有 IMG_PH 预览框段后保存
d2 = Document(str(NEW))
removed = 0
for p in list(d2.paragraphs):
    descrs = [el.get("descr") or "" for el in p._p.iter() if el.tag.endswith("}docPr")]
    if any(x.startswith("IMG_PH:") for x in descrs):
        p._p.getparent().remove(p._p)
        removed += 1
nopv = TMP / "new_no_preview_资格证明.docx"
d2.save(str(nopv))
print("previews removed:", removed)

# Word 导出三个文件
import win32com.client  # noqa: E402
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
word.DisplayAlerts = 0
try:
    for f in sorted(TMP.glob("*.docx")):
        pdf = TMP / (f.stem + ".pdf")
        try:
            doc = word.Documents.Open(str(f), ReadOnly=True)
            doc.ExportAsFixedFormat(str(pdf), 17)
            doc.Close(False)
            print("OK  ", f.name)
        except Exception as e:
            print("FAIL", f.name, "->", (getattr(e, "excepinfo", None) or [""])[2])
finally:
    word.Quit()
