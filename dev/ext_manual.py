# -*- coding: utf-8 -*-
"""提取人工版 docx 文本（段落+表格），供通道B 语义审查。"""
import sys, io
from pathlib import Path
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

SRC = Path(r"E:\bidcraft_skill\dev\manual_conv")
OUT = Path(r"E:\bidcraft_skill\dev\manual_txt")
OUT.mkdir(exist_ok=True)

files = sorted(SRC.glob("*.docx"))
for f in files:
    try:
        doc = Document(str(f))
    except Exception as e:
        print(f"== {f.name}: 打开失败 {e}")
        continue
    lines = []
    # 段落
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            lines.append(t)
    # 表格
    for ti, tbl in enumerate(doc.tables):
        lines.append(f"[表格{ti+1}]")
        for row in tbl.rows:
            cells = [c.text.strip().replace("\n", "/") for c in row.cells]
            # 去重相邻（合并单元格）
            out = []
            for c in cells:
                if not out or out[-1] != c:
                    out.append(c)
            lines.append(" | ".join(out))
    txt = "\n".join(lines)
    outp = OUT / (f.stem + ".txt")
    outp.write_text(txt, encoding="utf-8")
    print(f"{f.name}: {len(lines)} 行 / {len(txt)} 字符 -> {outp.name}")
