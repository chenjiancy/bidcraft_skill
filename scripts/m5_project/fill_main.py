# -*- coding: utf-8 -*-
"""⑦ filler 拆包 · 主流程（fill_project / 报告输出）。"""
import json
from pathlib import Path

from _shared import core

from .fill_common import FILL_ID, FILL_VERSION, HAVE_DOCX, FillError
from .fill_images import _fill_cell_image, _fill_image_placeholders
from .fill_tables import (_fill_performance_tables, _fill_personnel_tables,
                          _fill_resume_tables)
from .fill_text import _fill_text_placeholders, _scan_remaining_placeholders, _fmt_date

__all__ = ["_load_material", "_load_deadline", "fill_project", "_write_reports"]


def _load_material(proj_dir):
    p = Path(proj_dir) / "招标解析" / "素材清单.json"
    return core.read_json(str(p), None) or {}


def _load_deadline(proj_dir):
    p = Path(proj_dir) / "招标解析" / "投标要点.json"
    data = core.read_json(str(p), None) or {}
    return data.get("overview", {}).get("deadline", "") or \
        data.get("deadline", "") or "2026-10-09"


def fill_project(ent, proj_dir, out_dir=None, lib_root=None, register=True):
    """项目模板（冻结版）→ 商务标。返回统计。"""
    if not HAVE_DOCX:
        raise FillError("填充引擎依赖 python-docx，当前环境未安装")
    from docx import Document
    ent = Path(ent)
    proj_dir = Path(proj_dir)
    lib_root = Path(lib_root) if lib_root else ent / "企业级" / "素材库"
    tpl_dir = proj_dir / "项目模板"
    out = Path(out_dir) if out_dir else proj_dir / "商务标"
    out.mkdir(parents=True, exist_ok=True)

    material = _load_material(proj_dir)
    biz = material.get("企业基础信息", {}) or {}
    deadline = material.get("投标截止日期", "") or _load_deadline(proj_dir)

    stats = {"文件": [], "缺图": [], "待补": [], "清单外": [],
             "文字占位填充": 0, "占位框移除": 0}
    for f in sorted(tpl_dir.glob("*.docx")):
        doc = Document(str(f))
        missing, pending = [], []
        filled = _fill_text_placeholders(doc, material, biz, deadline, pending)
        n_prev, outside = _fill_image_placeholders(doc, lib_root, proj_dir, material, missing)
        stats["占位框移除"] += n_prev
        stats["清单外"] += outside
        _fill_personnel_tables(doc, material, pending)
        _fill_resume_tables(doc, lib_root, material, pending)
        _fill_performance_tables(doc, material, pending)
        # 残留占位符扫描 → 待补清单
        _scan_remaining_placeholders(doc, f.name, pending)
        # 表内图片占位（组织机构框图）
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    ct = cell.text.strip()
                    if ct.startswith("【图片：") and "组织机构框图" in ct:
                        src = Path(lib_root) / "企业介绍/组织机构图_20261001.png"
                        if src.is_file():
                            _fill_cell_image(cell, str(src), 15, 8.5)
                        else:
                            missing.append((ct, "组织机构图素材缺失"))
        doc.save(str(out / f.name))
        stats["文件"].append({"文件": f.name, "填充文字": filled,
                              "缺图": [x[0] for x in missing],
                              "待补": list(dict.fromkeys(pending))})
        stats["缺图"] += [tuple(x) for x in missing]
        stats["待补"] += pending
        stats["文字占位填充"] += filled

    stats["缺图"] = list(dict.fromkeys(stats["缺图"]))
    stats["待补"] = list(dict.fromkeys(stats["待补"]))
    stats["清单外"] = list(dict.fromkeys(stats["清单外"]))
    _write_reports(out, stats, deadline)
    if register:
        core.write_json(str(out / "生成记录.json"), {
            "id": FILL_ID, "version": FILL_VERSION, "日期": deadline,
            "项目": material.get("project", ""),
            "统计": {k: v for k, v in stats.items() if k not in ("文件",)},
        })
    return stats


def _write_reports(out, stats, deadline):
    lines = ["# 缺图清单（商务标生成 %s）" % _fmt_date(deadline), "",
             "以下图片素材缺失，占位段已删除；补齐后重新生成商务标：", ""]
    for ph, note in stats["缺图"]:
        lines.append("- %s（%s）" % (ph, note))
    if not stats["缺图"]:
        lines.append("- 无")
    (out / "缺图清单.md").write_text("\n".join(lines), encoding="utf-8")

    lines = ["# 清单外素材（商务标生成 %s）" % _fmt_date(deadline), "",
             "以下素材库文件**未在素材清单列明**，按白名单规则未插入商务标；"
             "如确需使用，请先在素材清单登记（propose/apply）后再重新生成：", ""]
    for rel in stats["清单外"]:
        lines.append("- %s" % rel)
    if not stats["清单外"]:
        lines.append("- 无")
    (out / "清单外素材.md").write_text("\n".join(lines), encoding="utf-8")

    lines = ["# 待补字段清单（商务标生成 %s）" % _fmt_date(deadline), "",
             "以下文字占位无值或待确认，**保留占位符**，交标前人工填写：", ""]
    for x in stats["待补"]:
        lines.append("- %s" % x)
    if not stats["待补"]:
        lines.append("- 无")
    (out / "待补字段清单.md").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="M5 填充引擎：项目模板 → 商务标")
    ap.add_argument("--ent", default=r"E:\监理标书制作\示例建设工程监理有限公司")
    ap.add_argument("--proj", default=r"E:\监理标书制作\示例建设工程监理有限公司\项目级\示例示例园区尾水水质提升工程（EPC总承包）监理")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    s = fill_project(args.ent, args.proj, out_dir=args.out)
    print(json.dumps({k: v for k, v in s.items() if k != "文件"},
                     ensure_ascii=False, indent=1))
