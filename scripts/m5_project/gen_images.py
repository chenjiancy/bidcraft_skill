# -*- coding: utf-8 -*-
"""⑦ generator 拆包 · 图片占位（【图片：xxx】带边框占位段 + 灰底预览框）。

预览框（方案A：模板直接可见位置/尺寸/内容）→ ph_preview 模块。
"""
from . import ph_preview as phprev   # 图片占位预览框

__all__ = [
    "_make_image_ph_para", "_set_picture_marker", "_make_ph_preview_para",
    "_insert_ph_preview_after", "_insert_image_ph_after_table",
    "_insert_image_ph_after_para",
]


def _make_image_ph_para(text, font):
    """构造【图片：xxx】占位段（模板库同款样式：左/下/右边框 + F2F2F2 底纹 + 加粗 28号）。"""
    from .gen_common import qn
    from docx.oxml import OxmlElement
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pBdr = OxmlElement("w:pBdr")
    for edge in ("w:left", "w:bottom", "w:right"):
        el = OxmlElement(edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), "808080")
        pBdr.append(el)
    pPr.append(pBdr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    p.append(pPr)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), font)
    rf.set(qn("w:hAnsi"), font)
    rf.set(qn("w:eastAsia"), font)
    rPr.append(rf)
    for tag in ("w:b", "w:bCs"):
        rPr.append(OxmlElement(tag))
    col = OxmlElement("w:color")
    col.set(qn("w:val"), "000000")
    rPr.append(col)
    for tag, val in (("w:sz", "28"), ("w:szCs", "28")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        rPr.append(el)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    p.append(r)
    return p


def _set_picture_marker(run, marker):
    """给段落内图片打标记：wp:docPr@descr = marker（填充引擎按此定位/删除预览框）。"""
    for el in run._r.iter():
        if el.tag.endswith("}docPr"):
            el.set("descr", marker)
            return True
    return False


def _make_ph_preview_para(doc, ph_text, assets=None):
    """构造图片占位预览框段：灰底占位图（尺寸=image_spec 目标口径），docPr 打标 IMG_PH:xxx。
    assets=[(素材绝对路径, 口径key, 是否换页), ...] → 框内直接画真实素材缩略图。"""
    from docx.shared import Cm
    png = phprev.make_placeholder_png(ph_text, assets=assets)
    w_cm, h_cm = phprev.box_size_for(ph_text)
    p = doc.add_paragraph()
    p.alignment = 1                                     # CENTER
    run = p.add_run()
    run.add_picture(png, width=Cm(w_cm), height=Cm(h_cm))
    _set_picture_marker(run, phprev.PH_PREVIEW_PREFIX + ph_text)
    return p


def _insert_ph_preview_after(doc, anchor_el, ph_text, assets=None):
    """把预览框段移动到锚点元素（占位文字段）之后；返回是否成功。
    L4 修复（2026-10-18）：锚点可能是 lxml 元素（资格表 _make_image_ph_para）或
    python-docx Paragraph（F16 独立文件 doc.add_paragraph）——统一取 _p 再 addnext。"""
    try:
        p = _make_ph_preview_para(doc, ph_text, assets=assets)
        anchor = getattr(anchor_el, "_p", anchor_el)
        anchor.addnext(p._p)
        return True
    except Exception:
        return False


def _insert_image_ph_after_table(doc, ctx_map, rules, font, assets_map=None):
    """在指定表格（按表前文关键词，去空格匹配）之后插入图片占位段 + 预览框。
    返回 (占位段数, 预览框数)。
    L4 修复（2026-10-18）：
      - 锚点匹配表格实例数 == 该组占位数（且 >1）→ 按实例逐表分配（附表8 简历表
        每人一份：第 i 份简历表后插第 i 人员图组，对齐人工标书 P35-58）；
      - 否则整组插第一个匹配表格后（原 L3 break 行为，防锚点多命中重复）。"""
    n = n_prev = 0
    assets_map = assets_map or {}
    for kw, phs in rules:
        instances = [t for t in doc.tables
                     if kw in ctx_map.get(t._tbl, "").replace(" ", "").replace("\u3000", "")]
        if not instances:
            continue
        if len(phs) > 1 and len(instances) == len(phs):
            # 实例分配：第 i 个表格后插 phs[i]
            for tbl, text in zip(instances, phs):
                el = _make_image_ph_para(text, font)
                tbl._tbl.addnext(el)
                n += 1
                if _insert_ph_preview_after(doc, el, text, assets_map.get(text)):
                    n_prev += 1
        else:
            # 整组插第一个匹配表格后
            prev = instances[0]._tbl
            for text in phs:
                el = _make_image_ph_para(text, font)
                prev.addnext(el)          # 正序插入
                prev = el
                n += 1
                if _insert_ph_preview_after(doc, el, text, assets_map.get(text)):
                    n_prev += 1
    return n, n_prev


def _insert_image_ph_after_para(doc, rules, font, assets_map=None):
    """在锚点段落（含关键词）之后插入图片占位段 + 预览框。
    返回 (占位段数, 预览框数)。
    L3 修复（2026-10-18）：每组规则只对**第一个匹配段落**插入一次——
    原实现对每个匹配段都插一组，锚点关键词（如「具体按招标公告…」）命中
    多个段落时占位重复（实测 12→25 处）。"""
    n = n_prev = 0
    assets_map = assets_map or {}
    for kw, phs in rules:
        for p in doc.paragraphs:
            if kw in p.text:
                prev = p._p
                for text in phs:
                    el = _make_image_ph_para(text, font)
                    prev.addnext(el)          # 正序插入
                    prev = el
                    n += 1
                    if _insert_ph_preview_after(doc, el, text, assets_map.get(text)):
                        n_prev += 1
                break                          # 只插第一处匹配
    return n, n_prev
