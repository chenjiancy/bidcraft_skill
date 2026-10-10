# -*- coding: utf-8 -*-
"""⑦ generator · 页数校验 + 独占页排版收敛（M5 重设计 v2.1，2026-10-10）。

用户规则（页面布局硬约束，2026-10-10 澄清）：
- 企业模板独占 1 页的内容（封面/法定代表人身份证明/授权委托书/简历表/图片/表格等）
  → 项目模板**必须独占 1 页，严禁跨页**，基本页面布局与企业模板一致；
- **收敛手段（用户明确授权）**：文字可删除空行、调整行距；表格可删除空白行、
  调整行高；不得影响阅读与页面美观；
- **不改变字号、不改变页边距**（保持企业模板格式契约值）。

实现（代码确定性）：
1. Word COM 全量页数校验（verify_pages.ps1，按模板选择映射配对）；
2. 独占页跨页文件 → 自动排版收敛（仅：删空行/调行距/删表格空白行/调行高/收段距）
   复核页数；仍跨页 → 报"待人工调整"（agent 检查兜底）。
3. 收敛过文件的基线重新登记（版本 bump，可追溯）。
"""
import subprocess
from pathlib import Path

from _shared import core
from .gen_common import HAVE_DOCX, GenError, Document, qn

__all__ = ["verify_and_fit", "page_count", "shrink_to_fit"]

# 收敛档位（顺序执行；文字一字不改，不动字号、不动页边距）——对齐用户 V5 确认顺序：
#   round0: 删空行 + 固定行距 exactly 28pt（正文基准）
#   round1+: 固定行距 24→22→20→18→16→14→12→11→10.5pt（低于字号自动取字号，防截字）
#   表格行高/单元格内边距随轮次同步收敛（round6 后固定 0.6cm/0.05cm 下限）
_LINE_STEPS = [("exact", 28.0), ("exact", 24.0), ("exact", 22.0), ("exact", 20.0),
               ("exact", 18.0), ("exact", 16.0), ("exact", 14.0), ("exact", 12.0),
               ("exact", 11.0), ("exact", 10.5), ("exact", 10.5), ("exact", 10.5)]
_SPACE_STEPS = [8.0, 6.0, 4.0, 2.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0]   # pt；段前段后收敛上限
_ROW_HEIGHT_STEPS = [1.5, 1.2, 1.0, 0.9, 0.8, 0.7, 0.6, 0.6, 0.6, 0.6, 0.6, 0.6]  # cm；显式行高上限
_CELL_MARGIN_STEPS = [None, 0.15, 0.10, 0.08, 0.06, 0.05, 0.05, 0.05, 0.05, 0.05, 0.05, 0.05]  # cm；单元格内边距上限

_PS1_DIR = Path(__file__).parent


def page_count(docx_path, timeout=90):
    """单文件页数（Word COM，verify_pages_one.ps1）。失败抛 GenError。"""
    ps1 = _PS1_DIR / "verify_pages_one.ps1"
    try:
        r = subprocess.run(
            ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass",
             "-File", str(ps1), "-File", str(Path(docx_path).resolve())],
            capture_output=True, text=True, encoding="utf-8", timeout=timeout)
    except Exception as ex:
        raise GenError("页数查询失败（Word COM 不可用？）：%s" % ex)
    if r.returncode != 0:
        raise GenError("页数查询失败：%s" % (r.stderr or r.stdout or "")[:200])
    try:
        return int(r.stdout.strip())
    except Exception:
        raise GenError("页数查询输出异常：%r" % r.stdout[:100])


def _para_is_blank(p):
    """段落是否空白（无文字、无图片、无书签/分隔符之外的内容）。"""
    if "".join(r.text for r in p.runs).strip():
        return False
    # 含图片/图形/文本框 → 不视为空白
    xml = p._p.xml
    if "graphic" in xml or "pict" in xml or "txbxContent" in xml:
        return False
    return True


def _row_is_blank(row):
    """表格行是否整行空白（所有单元格无文字）。"""
    seen = set()
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        if (c.text or "").strip():
            return False
    return True


def _apply_round(doc, round_i, delete_page_breaks=False):
    """应用第 round_i 轮排版收敛（删空行/行距/段距/表格空白行与行高/单元格内边距；
    不动字号与页边距）。delete_page_breaks=True 时删除分页符（独占页目标文件：
    企业模板无分页符、独占 1 页从文档起始页开始；原文拷贝可能带入分页符导致跨页）。
    返回是否实际改动。"""
    from docx.shared import Cm, Pt
    line_mode, line_val = _LINE_STEPS[min(round_i, len(_LINE_STEPS) - 1)]
    space = _SPACE_STEPS[min(round_i, len(_SPACE_STEPS) - 1)]
    row_h = _ROW_HEIGHT_STEPS[min(round_i, len(_ROW_HEIGHT_STEPS) - 1)]
    cell_m = _CELL_MARGIN_STEPS[min(round_i, len(_CELL_MARGIN_STEPS) - 1)]
    changed = False

    if delete_page_breaks:
        for br in doc.element.body.findall(".//" + qn("w:br")):
            if br.get(qn("w:type")) == "page":
                parent = br.getparent()
                parent.remove(br)
                changed = True

    def fit_line(pf, max_pt):
        """行距收敛：mult 档（倍数）或 exact 档（固定值，不低于文档最大字号防截字）。"""
        nonlocal changed
        if line_mode == "mult":
            if pf.line_spacing and pf.line_spacing > line_val:
                pf.line_spacing = line_val
                changed = True
        else:
            cur = pf.line_spacing
            target = max(line_val, max_pt)      # 固定值不得低于字号（防截字）
            if cur is None or cur != target:
                pf.line_spacing = Pt(target)
                changed = True

    def max_font_pt(paras):
        m = 0.0
        for p in paras:
            for r in p.runs:
                if r.font.size is not None:
                    m = max(m, r.font.size.pt)
        return m

    def shrink_paras(paras):
        nonlocal changed
        mpt = max_font_pt(paras)
        for p in paras:
            if _para_is_blank(p):
                el = p._p
                el.getparent().remove(el)
                changed = True
                continue
            pf = p.paragraph_format
            fit_line(pf, mpt)
            if pf.space_before and pf.space_before.pt > space:
                pf.space_before = Pt(space)
                changed = True
            if pf.space_after and pf.space_after.pt > space:
                pf.space_after = Pt(space)
                changed = True

    # 正文顶层空段落删除 + 行距/段距收敛
    body_paras = list(doc.paragraphs)
    mpt = max_font_pt(body_paras)
    for p in body_paras:
        if _para_is_blank(p):
            el = p._p
            el.getparent().remove(el)
            changed = True
        else:
            pf = p.paragraph_format
            fit_line(pf, mpt)
            if pf.space_before and pf.space_before.pt > space:
                pf.space_before = Pt(space)
                changed = True
            if pf.space_after and pf.space_after.pt > space:
                pf.space_after = Pt(space)
                changed = True

    # 表格：空白行删除 + 单元格内空段收敛 + 行高收敛 + 单元格内边距收敛
    for table in doc.tables:
        blank_rows = [row for row in table.rows if _row_is_blank(row)]
        for row in blank_rows:
            tr = row._tr
            tr.getparent().remove(tr)
            changed = True
        for row in table.rows:
            # 行高（仅显式设置的较大行高下调；auto/无设置不动）
            trPr = row._tr.trPr
            if trPr is not None:
                for th in trPr.findall(qn("w:trHeight")):
                    try:
                        v = int(th.get(qn("w:val")) or 0)   # twips
                        if v > int(row_h * 567):            # 1cm ≈ 567 twips
                            th.set(qn("w:val"), str(int(row_h * 567)))
                            changed = True
                    except Exception:
                        pass
            cell_paras = []
            for cell in row.cells:
                cell_paras.extend(cell.paragraphs)
            mpt_c = max_font_pt(cell_paras)
            for p in cell_paras:
                if _para_is_blank(p):
                    el = p._p
                    el.getparent().remove(el)
                    changed = True
                else:
                    pf = p.paragraph_format
                    fit_line(pf, mpt_c)
                    if pf.space_before and pf.space_before.pt > space:
                        pf.space_before = Pt(space)
                        changed = True
                    if pf.space_after and pf.space_after.pt > space:
                        pf.space_after = Pt(space)
                        changed = True
        # 单元格内边距（tblCellMar）收敛
        if cell_m is not None:
            tblPr = table._tbl.tblPr
            if tblPr is not None:
                mar = tblPr.find(qn("w:tblCellMar"))
                if mar is None:
                    mar = tblPr.makeelement(qn("w:tblCellMar"), {})
                    tblPr.append(mar)
                for side in ("w:top", "w:left", "w:bottom", "w:right"):
                    el = mar.find(qn(side))
                    if el is None:
                        el = mar.makeelement(qn(side), {})
                        mar.append(el)
                    try:
                        v = int(el.get(qn("w:w")) or 0)     # twips
                        if v > int(cell_m * 567):
                            el.set(qn("w:w"), str(int(cell_m * 567)))
                            el.set(qn("w:type"), "dxa")
                            changed = True
                    except Exception:
                        pass
    return changed


def shrink_to_fit(docx_path, timeout=90):
    """独占页文件排版收敛：删空行/调行距/表格删空白行调行高（不缩字号、不动页边距），
    然后复核一次页数。不再逐轮 Word COM 查页数（性能：每次查页数需重启 Word，太慢）。
    返回 {"达成": bool, "收敛手段": str, "页数": p}。"""
    if not HAVE_DOCX:
        raise GenError("排版收敛依赖 python-docx")
    p = Path(docx_path)
    for i in range(len(_LINE_STEPS)):
        doc = Document(str(p))
        _apply_round(doc, i, delete_page_breaks=True)
        doc.save(p)
    try:
        pages = page_count(p, timeout=timeout)
    except GenError as ex:
        return {"达成": False, "收敛手段": "空行删除/行距/表格行收敛", "页数": -1, "错误": str(ex)}
    return {"达成": pages == 1, "收敛手段": "空行删除/行距/表格行收敛", "页数": pages}


def verify_and_fit(ent, project, out_dir, format_config, results,
                   register_baseline=True, timeout=90):
    """页数校验 + 独占面排版收敛（proj-gen 内调用，V5：格式配置独占一面清单）。

    返回 {"结论", "文件":[{文件, 独占一面要求, 项目模板页数, 结论, 说明, 压缩}]}。
    """
    from m_feedback import feedback as fb
    from .format_applier import load_format_config
    out = Path(out_dir)
    ps1 = _PS1_DIR / "verify_pages.ps1"
    cfg = format_config or load_format_config()
    # 文件级独占（纯文件名）清单：封面/密封袋封面/开标/附录/法代/授权/中小企业声明函/承诺书
    own_files = [f for f in cfg.get("独占一面", {}).get("文件级（每个文件独立成页）", [])
                 if "（" not in f]
    if not ps1.is_file():
        raise GenError("页数校验缺少脚本：%s" % ps1)
    cfg_json = _format_config_path(format_config)
    r = subprocess.run(
        ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass",
         "-File", str(ps1), "-FormatConfigJson", cfg_json,
         "-ProjDir", str(out.resolve()), "-OutJson", str((out / "页数校验.json").resolve())],
        capture_output=True, text=True, encoding="utf-8", timeout=timeout * 2)
    if r.returncode != 0:
        raise GenError("页数校验失败（Word COM 不可用？）：%s" % (r.stderr or r.stdout or "")[:300])
    items = []
    if (out / "页数校验.json").is_file():
        items = core.read_json(out / "页数校验.json", []) or []
    problems = []
    for it in items:
        fname = it.get("文件", "")
        own = it.get("独占一面要求", False)
        p_p = it.get("项目模板页数", 0)
        conclusion = it.get("结论", "跳过")
        note = it.get("说明", "")
        shrink = None
        if conclusion == "不通过" and own and p_p != 1:
            # 独占面跨页 → 自动排版收敛（文字一字不改，按格式配置收敛顺序）
            try:
                shrink = shrink_to_fit(out / fname, timeout=timeout)
                if shrink.get("达成"):
                    conclusion = "通过（已收敛）"
                    note = "格式配置要求独占一面；排版收敛至 1 页（%s）" % shrink.get("收敛手段", "空行/行距/表格行")
                    if register_baseline:
                        try:
                            rel = "项目级/%s/项目模板/%s" % (project, fname)
                            fb.register(ent, rel, level="项目级", project=project,
                                        ptype="项目模板", generator="m5-project-gen",
                                        note="独占面排版收敛（%s）" % shrink.get("收敛手段", ""))
                        except Exception as ex:
                            note += "；基线更新失败：%s" % ex
                else:
                    conclusion = "不通过"
                    note += "；排版收敛后仍 %s 页（需人工调整排版或 agent 检查复核）" % shrink.get("页数", "?")
            except Exception as ex:
                conclusion = "不通过"
                note += "；收敛失败：%s" % ex
        if conclusion == "不通过":
            problems.append({"文件": fname, "说明": note})
        it["独占一面要求"] = own
        it["项目模板页数"] = p_p
        it["结论"] = conclusion
        it["说明"] = note
        it["压缩"] = shrink
    allok = all(it.get("结论", "跳过") in ("通过", "通过（已收敛）", "跳过") for it in items)
    return {"结论": "通过" if allok else "不通过", "文件": items, "问题": problems}


def _format_config_path(p):
    if p:
        return str(Path(p))
    return str(Path(__file__).resolve().parent / "format_config.json")
