# -*- coding: utf-8 -*-
"""
M2 企业级模板库占位预览图集（2026-10-04 用户确认）。

目标：企业级模板库范本（模板库/<代理机构>/<方式>/*.docx）不改原文件，
另出一份「占位预览图集」——每个【图片：xxx】占位段后插入预览框
（复用 M5 预览框 v2）：
  - 单图素材（FILL_IMG_MAP 固定路径 或 关键词兜底命中企业级素材库）→ 框内直接显示
    将填充的真实素材缩略图 + 文件名 + 尺寸；
  - 组合/项目相关占位（业绩材料、人员证书组、荣誉等，素材随项目而定）→ 灰底
    【待补素材】标注（按项目素材清单填充），人工在模板阶段即可看到占位位置与大小。

原文件只读（python-docx 打开后另存副本），输出到 企业级/模板库/_占位预览图集/
<代理机构>/<方式>/，并生成 占位预览图集说明.md 汇总每文件占位与素材状态。
"""
from pathlib import Path

try:
    from docx import Document
    HAVE_DOCX = True
except Exception:
    HAVE_DOCX = False

from _shared import docx_util              # noqa: E402
from m5_project import image_spec as imgsp  # noqa: E402
from m5_project import ph_preview as phprev  # noqa: E402
from m5_project.fill_images import FILL_IMG_MAP  # noqa: E402
from m5_project.gen_images import _set_picture_marker  # noqa: E402

# 关键词 → (素材相对路径, 口径键)：FILL_IMG_MAP 未精确命中时的企业级素材库兜底
_KEYWORD_ASSETS = [
    ("开户许可证", ("财务/财务证照/开户许可证_长期.png", "开户许可证扫描件")),
    ("营业执照", ("资质/营业执照_副本_长期.png", "营业执照扫描件")),
    ("组织机构", ("企业介绍/组织机构图_20261001.png", "组织机构框图")),
]

# 关键词 → 口径键（预览框尺寸兜底：组合/项目相关占位也按真实填充口径显示大小）
_KEYWORD_SPEC = [
    ("三体系", "三体系认证证书"),
    ("职称证书", "其他监理人员职称证书"),
    ("监理示范", "监理示范优质工程"),
    ("先进", "先进优秀监理企业证书"),
    ("身份证", "身份证"),
]

DEFAULT_OUT_DIR_NAME = "_占位预览图集"


def _resolve_assets(ph_text, lib_root):
    """占位文案 → 素材列表；None = 待补（组合/项目相关，按项目素材清单填充）。"""
    mapped = FILL_IMG_MAP.get(ph_text)
    if mapped and mapped[0]:
        full = Path(lib_root) / mapped[0]
        return [(str(full), mapped[1], True)] if full.is_file() else None
    for kw, (rel, skey) in _KEYWORD_ASSETS:
        if kw in ph_text:
            full = Path(lib_root) / rel
            return [(str(full), skey, True)] if full.is_file() else None
    return None


def _preview_box_size(ph_text):
    """企业级预览框尺寸：精确键 → image_spec；关键词兜底；否则默认代表框。"""
    key = imgsp.spec_for(ph_text)
    if key and key in imgsp.IMG_SPEC:
        return phprev.box_size_for(ph_text)
    for kw, skey in _KEYWORD_SPEC:
        if kw in ph_text and skey in imgsp.IMG_SPEC:
            w = float(imgsp.IMG_SPEC[skey].get("宽", 16))
            h = imgsp.IMG_SPEC[skey].get("高", None)
            return (w, float(h)) if isinstance(h, (int, float)) else (w, 8.0)
    return phprev.box_size_for(ph_text)


def _make_preview_para(doc, ph_text, assets, size):
    """构造预览框段（v2）：真实素材缩略图或【待补素材】灰底，docPr 打标 IMG_PH:xxx。"""
    from docx.shared import Cm
    png = phprev.make_placeholder_png(ph_text, assets=assets, box_size=size)
    w_cm, h_cm = size
    p = doc.add_paragraph()
    p.alignment = 1                                     # CENTER
    run = p.add_run()
    run.add_picture(png, width=Cm(w_cm), height=Cm(h_cm))
    _set_picture_marker(run, phprev.PH_PREVIEW_PREFIX + ph_text)
    return p


def build_template_previews(ent, agency, mode, out_dir=None, lib_root=None):
    """为企业级模板库范本生成占位预览副本（不修改原文件）。

    返回 {"输出目录", "文件": [{文件, 图片占位, 预览框, 素材数, 待补数}], "说明文件"}。
    """
    if not HAVE_DOCX:
        from _shared.core import LibraryError
        raise LibraryError("预览图集依赖 python-docx，当前环境未安装")
    ent = Path(ent)
    root = ent / "企业级" / "模板库" / agency / mode
    if not root.is_dir():
        from _shared.core import LibraryError
        raise LibraryError("模板库目录不存在：%s" % root)
    lib = Path(lib_root) if lib_root else ent / "企业级" / "素材库"
    out = Path(out_dir) if out_dir else ent / "企业级" / "模板库" / DEFAULT_OUT_DIR_NAME / agency / mode
    out.mkdir(parents=True, exist_ok=True)

    records = []
    total_ph = total_prev = total_assets = total_pending = 0
    for src in sorted(root.glob("*.docx")):
        doc = Document(str(src))
        found = []                                    # [(占位文案, 是否有素材)]
        for p in list(doc.paragraphs):
            t = docx_util.para_text(p).strip()
            if not (t.startswith("【图片：") and t.endswith("】")):
                continue
            assets = _resolve_assets(t, lib)
            size = _preview_box_size(t)
            prev = _make_preview_para(doc, t, assets, size)
            p._p.addnext(prev._p)
            found.append((t, assets is not None))
        dst = out / ("预览_" + src.name)
        doc.save(str(dst))
        n_assets = sum(1 for _, ok in found if ok)
        n_pending = len(found) - n_assets
        records.append({"文件": src.name, "图片占位": len(found), "预览框": len(found),
                        "素材数": n_assets, "待补数": n_pending})
        total_ph += len(found)
        total_prev += len(found)
        total_assets += n_assets
        total_pending += n_pending

    manifest = _write_manifest(out, agency, mode, records,
                               total_ph, total_prev, total_assets, total_pending)
    return {"输出目录": str(out), "文件": records, "说明文件": str(manifest),
            "合计": {"图片占位": total_ph, "预览框": total_prev,
                    "有素材": total_assets, "待补": total_pending}}


def _write_manifest(out_dir, agency, mode, records, total_ph, total_prev,
                    total_assets, total_pending):
    """汇总说明.md：每文件占位清单 + 素材状态。"""
    from _shared import core
    lines = ["# 企业级模板库 · 占位预览图集说明", ""]
    lines.append("- 模板库：%s ｜ %s" % (agency, mode))
    lines.append("- 生成时间：%s" % core.now_iso())
    lines.append("- 用途：原文件只读；本图集在副本中为每个【图片：xxx】占位插入预览框")
    lines.append("  （框内显示将填充的真实素材缩略图+文件名+尺寸；组合/项目相关占位")
    lines.append("  显示【待补素材】灰底框，素材随项目素材清单而定）。")
    lines.append("- 说明：预览框见即可审；表格内图片占位（如附表1 组织机构框图）")
    lines.append("  暂不插框（会撑开表结构），按原路径在商务标填充阶段处理。")
    lines.append("")
    lines.append("合计：图片占位 %d ｜ 预览框 %d ｜ 有素材 %d ｜ 待补 %d"
                 % (total_ph, total_prev, total_assets, total_pending))
    lines.append("")
    for r in records:
        lines.append("## %s（图片占位 %d ｜ 预览框 %d ｜ 有素材 %d ｜ 待补 %d）"
                     % (r["文件"], r["图片占位"], r["预览框"], r["素材数"], r["待补数"]))
        doc = Document(str(out_dir / ("预览_" + r["文件"])))
        for p in doc.paragraphs:
            t = docx_util.para_text(p).strip()
            if t.startswith("【图片："):
                lines.append("- %s" % t)
        lines.append("")
    manifest = out_dir / "占位预览图集说明.md"
    manifest.write_text("\n".join(lines), encoding="utf-8")
    return str(manifest)
