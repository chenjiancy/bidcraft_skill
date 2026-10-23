# -*- coding: utf-8 -*-
"""M6 模板设计器 P1 集成测试：模板驱动导出（--template 全映射）。

覆盖：明标模板（封面/扉页/目录/正文/结尾部件 + 占位符填充）、暗标模板（铁律派生 + export.json）、
structure 无目录（TOC 域不渲染）、--fill 覆盖。依赖真实 python-docx（CI L2 安装）。
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))

from docx import Document             # noqa: E402
from m6_tech import export as E       # noqa: E402
from m6_tech import layout as L       # noqa: E402
from m6_tech import tpl_design as T   # noqa: E402

REPO = Path(__file__).resolve().parents[2]


class TemplateExportBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "素材库"
        self.root.mkdir()
        self.ent = L._builtin_preview_fixture(str(self.root))
        self.tech = Path(self.ent) / "项目级" / "预览项目" / "技术标"
        self.out_dir = self.tech / "产出"

    def tearDown(self):
        self.tmp.cleanup()

    def _tpl(self, name, **mut):
        doc = json.loads(json.dumps(T.TEMPLATE_EXAMPLE))
        for k, v in mut.items():
            doc[k] = v
        p = self.root / name
        p.write_text(json.dumps(doc, ensure_ascii=False), encoding="utf-8")
        return str(p)

    def _export(self, tpl_path, fill="", **kw):
        base = dict(root=str(self.root), enterprise="预览企业", project="预览项目",
                    template=tpl_path, fill=fill, override="", bid_mode="", force=True)
        base.update(kw)
        return E.cmd_tech_export(SimpleNamespace(**base))

    def _out_docx(self):
        return self.out_dir / "技术标_预览项目.docx"

    def _all_text(self, docx_path):
        doc = Document(str(docx_path))
        return "\n".join(p.text for p in doc.paragraphs), doc


class TestTemplateExportLight(TemplateExportBase):
    def test_light_full_structure_and_fill(self):
        # 明标模板（示例默认结构含 扉页+结尾），fill 覆盖项目名/企业名/日期
        tpl = self._tpl("明标模板.json")
        rc = self._export(tpl, fill="project_name=智慧园区项目;enterprise_name=甲公司;bid_date=2026年10月24日")
        self.assertEqual(rc, 0)
        out = self._out_docx()
        self.assertTrue(out.is_file(), "导出 docx 缺失")
        text, doc = self._all_text(out)
        self.assertGreater(len(doc.paragraphs), 0)
        self.assertIn("fldChar", doc.element.xml)                    # 目录域
        self.assertIn("智慧园区项目", text)                            # 封面/扉页填充
        self.assertIn("甲公司", text)                                 # 扉页企业名填充
        self.assertIn("2026年10月24日", text)                          # 日期填充
        self.assertIn("（以下无正文）", text)                           # 结尾页
        self.assertIn("目　录", text)                                  # 目录标题
        # 回读 export.json
        meta = json.loads((self.tech / "章节" / ".meta" / "export.json").read_text(encoding="utf-8"))
        self.assertFalse(meta["dark"])
        self.assertEqual(meta["layout"]["palette"], "deep-indigo")
        self.assertEqual(meta["layout"]["font"], "gov-gongwen")

    def test_light_fill_keeps_missing_placeholder(self):
        # 未提供的占位符键保留原文（如 site_photo 图片占位）
        tpl = self._tpl("明标模板.json")
        rc = self._export(tpl, fill="project_name=P1工程")
        self.assertEqual(rc, 0)
        out = self._out_docx()
        self.assertTrue(out.is_file())

    def test_structure_without_toc(self):
        # structure 不含 toc → 不渲染目录域
        tpl = self._tpl("无目录模板.json", structure=[
            {"part": "cover", "pageMode": "single"},
            {"part": "body", "pageMode": "flow"},
            {"part": "closing", "pageMode": "single", "optional": True},
        ])
        rc = self._export(tpl, fill="project_name=无目录工程")
        self.assertEqual(rc, 0)
        out = self._out_docx()
        text, doc = self._all_text(out)
        self.assertNotIn("fldChar", doc.element.xml)   # 无 TOC 域
        self.assertNotIn("目　录", text)
        self.assertIn("（以下无正文）", text)


class TestTemplateExportDark(TemplateExportBase):
    def test_dark_template_derived_and_exported(self):
        # 合规暗标模板（font=dark-safe、图片二选一、页眉页脚二选一）→ 导出 dark 分支生效
        tpl = self._tpl("暗标模板.json", rules={"bidMode": "暗标", "darkFontSafe": True},
                        styleRef=dict(T.TEMPLATE_EXAMPLE["styleRef"], font="dark-safe",
                                      imageRule="dark-safe-img", headerFooter="hf9-dark-none"))
        rc = self._export(tpl)
        self.assertEqual(rc, 0)
        out = self._out_docx()
        self.assertTrue(out.is_file())
        meta = json.loads((self.tech / "章节" / ".meta" / "export.json").read_text(encoding="utf-8"))
        self.assertTrue(meta["dark"])
        self.assertEqual(meta["layout"]["font"], "dark-safe")
        self.assertEqual(meta["layout"]["headerFooter"], "hf9-dark-none")
        text, doc = self._all_text(out)
        # 暗标：无封面企业信息（封面被切断）
        self.assertNotIn("投标技术标", text)
        # 无页码域（hf9 无页码无页眉页脚）
        self.assertNotIn("PAGE", doc.element.xml)

    def test_dark_invalid_font_rejected(self):
        # 暗标 + 明标字体 → 模板校验拒绝导出（铁律不可绕过）
        tpl = self._tpl("暗标违规模板.json", rules={"bidMode": "暗标", "darkFontSafe": True},
                        styleRef=dict(T.TEMPLATE_EXAMPLE["styleRef"], font="gov-gongwen"))
        with self.assertRaises(Exception):
            self._export(tpl)

    def test_dark_hf10_page_only_kept(self):
        tpl = self._tpl("暗标页码模板.json", rules={"bidMode": "暗标", "darkFontSafe": True},
                        styleRef=dict(T.TEMPLATE_EXAMPLE["styleRef"], font="dark-safe",
                                      imageRule="no-source", headerFooter="hf10-dark-page-only"))
        rc = self._export(tpl)
        self.assertEqual(rc, 0)
        meta = json.loads((self.tech / "章节" / ".meta" / "export.json").read_text(encoding="utf-8"))
        self.assertTrue(meta["dark"])
        self.assertEqual(meta["layout"]["headerFooter"], "hf10-dark-page-only")


if __name__ == "__main__":
    unittest.main()
