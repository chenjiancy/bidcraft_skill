# -*- coding: utf-8 -*-
"""M9 集成测试（L2）：真实项目全链体检 ——
半程（仅 M4+M6 docx）→ m9-verify 红灯；继续 check-rule→eval-run→remedy-gen→remedy-done→
修复→remedy-recheck 完整闭环 → m9-verify 绿灯 + m9-evidence 留痕包（KPI 提取正确）。
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(_HERE, "..", "..", "scripts"))

import bidcraft                       # noqa: E402
from docx.oxml import OxmlElement     # noqa: E402
from docx.oxml.ns import qn           # noqa: E402

PROJECT = "T19集成项目"
ENT = "T19集成企业"


def _points():
    return {"project": PROJECT, "modules": [
        {"id": "scoring", "title": "评分办法", "type": "mixed",
         "method": {"value": "综合评分法", "evidence": "e"},
         "score_items": [{"item": "监理大纲（50分）", "weight": 50, "unit": "分", "evidence": "E1"}]},
        {"id": "reject", "title": "废标红线", "type": "items",
         "items": [{"type": "格式", "rule": "技术标总页数不超过 60 页；按明标编制", "evidence": "e"}]},
    ]}


def _material(perf_path):
    return {"project": PROJECT,
            "personnel": [{"name": "赵监理", "role": "总监理工程师"}],
            "企业基础信息": {"法定代表人姓名": "钱总", "企业名称": ENT},
            "performance": [{"name": "T19集成业绩", "path": perf_path}],
            "qualification_required": [], "iso_certificates": [], "honors": [], "items": []}


def _mk_tech_docx(path, with_blank=False):
    from docx import Document
    doc = Document()
    doc.add_paragraph("封面：技术标")
    p = doc.add_paragraph()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    run = p.add_run()
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    p.add_run("目录")
    for i in range(8):
        if i == 2:
            para = doc.add_paragraph("第1章 监理工作内容")
            pPr = para._p.get_or_add_pPr()
            ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), "0"); pPr.append(ol)
        else:
            doc.add_paragraph("监理正文段落%d：全过程监理。" % i)
    if with_blank:
        doc.add_paragraph("")
        pb = doc.add_paragraph()
        br = OxmlElement("w:br"); br.set(qn("w:type"), "page")
        pb.add_run()._r.append(br)
        doc.add_paragraph("")
    doc.add_paragraph("本工程由总监理工程师赵监理负责，企业完成T19集成业绩项目。")
    doc.add_paragraph("结尾段：本技术标制作完成。")
    doc.save(str(path))


def _remove_blank_pages(path):
    from docx import Document
    doc = Document(str(path))
    paras = doc.paragraphs
    targets = []
    n = len(paras)
    for i, p in enumerate(paras):
        xml = p._p.xml
        if 'w:br' in xml and 'w:type="page"' in xml:
            lo = max(0, i - 1)
            hi = min(n - 1, i + 1)
            if all(not (paras[j].text or "").strip() for j in (lo, hi)):
                targets += [lo, i, hi]
    for idx in sorted(set(targets), reverse=True):
        paras[idx]._p.getparent().remove(paras[idx]._p)
    doc.save(str(path))


class M9VerifyIntegrationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._run("init-enterprise", "--name", ENT)
        self.ent = self.root / ENT
        self.proj = self.ent / "项目级" / PROJECT
        self.tdir = self.proj / "招标解析"
        self.tdir.mkdir(parents=True)
        (self.tdir / ("投标要点_%s.json" % PROJECT)).write_text(
            json.dumps(_points(), ensure_ascii=False), encoding="utf-8")
        perf = self.proj / "业绩" / "T19集成业绩.pdf"
        perf.parent.mkdir(parents=True)
        perf.write_bytes(b"%PDF-1.4 fake")
        (self.tdir / ("素材清单_%s.json" % PROJECT)).write_text(
            json.dumps(_material(str(perf)), ensure_ascii=False), encoding="utf-8")
        tech = self.proj / "技术标"
        tech.mkdir(parents=True)
        (tech / "设置.json").write_text(
            json.dumps({"schema": "m6.settings/v1", "bidMode": "明标",
                        "pageConstraint": {"max": 60}}, ensure_ascii=False), encoding="utf-8")
        out = tech / "产出"
        out.mkdir(parents=True)
        self.docx = out / ("技术标_%s.docx" % PROJECT)
        _mk_tech_docx(self.docx, with_blank=True)

    def tearDown(self):
        self.tmp.cleanup()

    def _run(self, *argv):
        import io as _io
        from contextlib import redirect_stdout, redirect_stderr
        out, err, code = _io.StringIO(), _io.StringIO(), 0
        with redirect_stdout(out), redirect_stderr(err):
            try:
                bidcraft.main(["--root", str(self.root), "--enterprise", ENT] + list(argv))
            except SystemExit as e:
                code = e.code or 0
        return code, out.getvalue(), err.getvalue()

    def _check_path(self):
        return self.proj / "标书检查" / ("机器检查_%s_技术标.json" % PROJECT)

    def _verify(self):
        code, out, err = self._run("m9-verify", "--project", PROJECT)
        self.assertEqual(code, 0, "%s|%s" % (out, err))
        vp = self.proj / "自检" / ("自检体检_%s.json" % PROJECT)
        return json.loads(vp.read_text(encoding="utf-8"))

    def test_half_and_full_loop(self):
        # —— 半程：仅 M4+M6 docx（未检查/未评标）→ 红灯 ——
        doc = self._verify()
        self.assertEqual(doc["总评"], "红灯")
        ids = {c["id"]: c["状态"] for c in doc["检查明细"]}
        self.assertEqual(ids["m7_check"], "fail")
        self.assertEqual(ids["m8_eval"], "fail")

        # —— 检查 + 评标 ——
        code, _, err = self._run("check-rule", "--project", PROJECT, "--out", str(self.docx), "--kind", "tech")
        self.assertEqual(code, 0, err)
        code, _, err = self._run("eval-run", "--project", PROJECT)
        self.assertEqual(code, 0, err)

        # —— 补救闭环：任务单 + 勾选 + 修复 + 重检 ——
        code, _, err = self._run("remedy-gen", "--project", PROJECT, "--check", str(self._check_path()))
        self.assertEqual(code, 0, err)
        tp = self.proj / "模拟评标" / ("失分补救任务单_%s.json" % PROJECT)
        tasks = json.loads(tp.read_text(encoding="utf-8"))["tasks"]
        blank = next(t for t in tasks if (t.get("检查快照") or {}).get("id") == "tech_blank_pages")
        code, _, err = self._run("remedy-done", "--project", PROJECT, "--id", blank["id"])
        self.assertEqual(code, 0, err)
        _remove_blank_pages(self.docx)
        code, _, err = self._run("check-rule", "--project", PROJECT, "--out", str(self.docx), "--kind", "tech")
        self.assertEqual(code, 0, err)
        code, _, err = self._run("remedy-recheck", "--project", PROJECT)
        self.assertEqual(code, 0, err)

        # —— 闭环后：绿灯 ——
        doc = self._verify()
        self.assertEqual(doc["总评"], "绿灯", [c for c in doc["检查明细"] if c["状态"] != "pass"])
        self.assertTrue(all(c["状态"] == "pass" for c in doc["检查明细"]),
                        [c for c in doc["检查明细"] if c["状态"] != "pass"])

        # —— 留痕包 ——
        code, out, err = self._run("m9-evidence", "--project", PROJECT)
        self.assertEqual(code, 0, "%s|%s" % (out, err))
        ep = self.proj / "自检" / ("投标留痕包_%s.json" % PROJECT)
        ev = json.loads(ep.read_text(encoding="utf-8"))
        self.assertEqual(ev["总评"], "绿灯")
        ratio = ev["KPI"]["检查项通过"].split("/")
        self.assertEqual(len(ratio), 2)
        self.assertEqual(ratio[0], ratio[1], "检查项应全部通过，实际 %s" % ev["KPI"]["检查项通过"])
        self.assertTrue(ev["KPI"]["任务完成率"].startswith("1/"))
        names = {it["产物"]: it for it in ev["产物清单"]}
        self.assertTrue(names["补救后重检报告"]["存在"])
        self.assertTrue((self.proj / "自检" / ("投标留痕包_%s.html" % PROJECT)).is_file())


if __name__ == "__main__":
    unittest.main()
