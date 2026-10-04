# -*- coding: utf-8 -*-
"""M5 填充引擎集成测试：项目模板 → 商务标（OCR 简历填充、业绩表、缺图清单、封面日期）。
依赖真实项目数据（E:\\监理标书制作\\示例建设工程监理有限公司）。"""
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))
from m5_project import filler as F                       # noqa: E402

ENT = r"E:\监理标书制作\示例建设工程监理有限公司"
PROJ = os.path.join(ENT, "项目级", "示例示例园区尾水水质提升工程（EPC总承包）监理")
LIB = os.path.join(ENT, "企业级", "素材库")
RESUME_PNG = os.path.join(LIB, "人员", "张三", "简历", "简历_20261001_P0.png")


class TestParseCert(unittest.TestCase):
    def test_parse_cert(self):
        self.assertEqual(F._parse_cert("注册证34008007（房建+市政公用，2028.1.23）"),
                         ("注册证", "34008007"))
        self.assertEqual(F._parse_cert("省监理工程师岗位证书皖监师2020000823"),
                         ("省监理工程师岗位证书", "2020000823"))
        self.assertEqual(F._parse_cert("岗位证书皖监员202300206"),
                         ("岗位证书", "202300206"))
        self.assertEqual(F._parse_cert(""), ("", ""))


class TestOcrResume(unittest.TestCase):
    def test_ocr_resume_fields(self):
        f = F.ocr_resume_fields(RESUME_PNG, cert_str="注册证34008007（房建+市政公用，2028.1.23）")
        self.assertEqual(f.get("人员姓名"), "张三")
        self.assertEqual(f.get("性别"), "男")
        self.assertEqual(f.get("出生年月"), "1983.9")
        self.assertEqual(f.get("学历"), "本科")
        self.assertEqual(f.get("身份证号"), "342626198309180031")
        self.assertEqual(f.get("证书编号"), "34008007")
        self.assertIn("相关专业经历", f)
        self.assertIn("主要经历", f)


class TestFillProject(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.out = os.path.join(tempfile.gettempdir(), "bid_test")
        if os.path.isdir(cls.out):
            shutil.rmtree(cls.out)
        cls.stats = F.fill_project(ENT, PROJ, out_dir=cls.out)

    def test_10_docx_generated(self):
        files = [f for f in os.listdir(self.out) if f.endswith(".docx")]
        self.assertEqual(len(files), 10)

    def test_cover_date_is_deadline(self):
        from docx import Document
        d = Document(os.path.join(self.out, "封面.docx"))
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("2026年10月9日", text)

    def test_resume_tables_filled(self):
        from docx import Document
        d = Document(os.path.join(self.out, "资格证明及辅助资料表.docx"))
        resumes = [t for t in d.tables
                   if "出生年月" in "".join(c.text for c in t.rows[0].cells)
                   and len(t.rows) >= 11]
        self.assertEqual(len(resumes), 6)
        r0 = "".join(c.text for c in resumes[0].rows[0].cells)
        self.assertIn("张三", r0)
        self.assertIn("1983.9", r0)
        self.assertIn("男", r0)
        r1 = "".join(c.text for c in resumes[0].rows[1].cells)
        self.assertIn("本科", r1)

    def test_performance_rows_filled(self):
        from docx import Document
        d = Document(os.path.join(self.out, "资格证明及辅助资料表.docx"))
        found = False
        for t in d.tables:
            first = "".join(c.text for c in t.rows[0].cells).replace(" ", "")
            if "项目名称" in first and "（万元）" not in first:
                r1 = "".join(c.text for c in t.rows[1].cells)
                self.assertIn("示例横江片区污水管网整治工程", r1)
                self.assertIn("14594.33", r1)
                found = True
                break
        self.assertTrue(found, "未找到附表2 已完成工程汇总表")

    def test_missing_image_report(self):
        # 社保已上传、赵六/赵七职称用户确认省略 → 缺图归零
        self.assertEqual(len(self.stats["缺图"]), 0, "缺图应为 0，实际 %s" % self.stats["缺图"])
        self.assertTrue(os.path.exists(os.path.join(self.out, "缺图清单.md")))
        self.assertTrue(os.path.exists(os.path.join(self.out, "待补字段清单.md")))
        self.assertTrue(os.path.exists(os.path.join(self.out, "生成记录.json")))

    def test_stats_schema(self):
        """方案A：统计含「占位框移除」（旧模板无预览框时为 0，键必须存在）。"""
        self.assertIn("占位框移除", self.stats)

    def test_combo_images_inserted(self):
        # 组合图片插入：资质/三体系/职称/荣誉/身份证正反/社保（多页 P0/P1 逐页）
        from docx import Document
        from docx.oxml.ns import qn
        d = Document(os.path.join(self.out, "资格证明及辅助资料表.docx"))
        n = len(d.element.body.findall(".//" + qn("w:drawing")))
        self.assertGreater(n, 60, "资格证明应插入组合+社保 66 图，实际 %d" % n)
        for fn, min_n in (("法定代表人身份证明.docx", 2), ("授权委托书.docx", 2)):
            dd = Document(os.path.join(self.out, fn))
            nn = len(dd.element.body.findall(".//" + qn("w:drawing")))
            self.assertGreaterEqual(nn, min_n, "%s 应插入身份证正反面" % fn)


class TestFillPreviewPipeline(unittest.TestCase):
    """方案A 全链路：新生成模板（含预览框）→ 填充 → 预览框清零、真图插入。"""

    @classmethod
    def setUpClass(cls):
        cls.tmp = os.path.join(tempfile.gettempdir(), "bid_m5_ph")
        if os.path.isdir(cls.tmp):
            shutil.rmtree(cls.tmp)
        os.makedirs(cls.tmp)
        sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))
        from m5_project import generator as G
        from m5_project import ph_preview as P
        cls.P = P
        cls.res = G.generate(Path(ENT), Path(PROJ).name, out_dir=Path(cls.tmp) / "项目模板",
                             register_baseline=False)

    def _marker_count(self, doc_path):
        from docx import Document
        d = Document(doc_path)
        return sum(1 for p in d.paragraphs for el in p._p.iter()
                   if el.tag.endswith("}docPr")
                   and (el.get("descr") or "").startswith(self.P.PH_PREVIEW_PREFIX))

    def test_template_has_preview_boxes(self):
        n = self._marker_count(Path(self.res["目录"]) / "资格证明及辅助资料表.docx")
        self.assertGreaterEqual(n, 6, "新生成模板应含≥6 个预览框，实际 %d" % n)

    def test_fill_removes_preview_and_inserts_images(self):
        import json
        from docx import Document
        from docx.oxml.ns import qn
        doc_path = Path(self.res["目录"]) / "资格证明及辅助资料表.docx"
        mat = json.loads((Path(PROJ) / "招标解析" / "素材清单.json").read_text(encoding="utf-8"))
        doc = Document(str(doc_path))
        missing = []
        removed, outside = F._fill_image_placeholders(doc, LIB, PROJ, mat, missing)
        self.assertGreater(removed, 0, "应移除预览框")
        # ③ 白名单：真实素材库全部在素材清单列明 → 清单外为空
        self.assertEqual(outside, [], "组合构建清单外素材应为空，实际: %s" % outside)
        # 剩余预览框应为 0
        left = sum(1 for p in doc.paragraphs for el in p._p.iter()
                   if el.tag.endswith("}docPr")
                   and (el.get("descr") or "").startswith(self.P.PH_PREVIEW_PREFIX))
        self.assertEqual(left, 0)
        # 真图已插入（drawing 数应远大于预览框数）
        nd = len(doc.element.body.findall(".//" + qn("w:drawing")))
        self.assertGreater(nd, 60, "填充后应插入组合+社保 66 图，实际 %d" % nd)


if __name__ == "__main__":
    unittest.main(verbosity=2)
