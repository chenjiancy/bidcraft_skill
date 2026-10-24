# -*- coding: utf-8 -*-
"""M10 T18 集成测试（L2）：补救→重检闭环真实项目链路 ——
合成含空白页的技术标 docx → check-rule（tech_blank_pages 失败）→ remedy-gen --check（生成整改任务）
→ remedy-done（勾选完成）→ 修复 docx（删空段）→ 重跑 check-rule → remedy-recheck（验证通过 + HTML 报告）。
未修复场景 → remedy-recheck 判"未通过"。
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(_HERE, "..", "..", "scripts"))

import bidcraft                       # noqa: E402
from docx.oxml import OxmlElement     # noqa: E402
from docx.oxml.ns import qn           # noqa: E402

PROJECT = "T18集成项目"
ENT = "T18集成企业"


def _points():
    return {"project": PROJECT, "modules": [
        {"id": "scoring", "title": "评分办法", "type": "mixed",
         "method": {"value": "综合评分法", "evidence": "e"},
         "score_items": [{"item": "监理大纲（50分）", "weight": 50, "unit": "分", "evidence": "E1"}]},
        {"id": "reject", "title": "废标红线", "type": "items",
         "items": [{"type": "格式", "rule": "技术标总页数不超过 60 页；按明标编制", "evidence": "e"}]},
    ]}


def _material(perf_path):
    return {"project": PROJECT,
            "personnel": [{"name": "赵监理", "role": "总监理工程师"}],
            "企业基础信息": {"法定代表人姓名": "钱总", "企业名称": ENT},
            "performance": [{"name": "T18集成业绩", "path": perf_path}],
            "qualification_required": [], "iso_certificates": [], "honors": [], "items": []}


def _mk_tech_docx(path, with_blank=False):
    from docx import Document
    doc = Document()
    doc.add_paragraph("封面：技术标")
    p = doc.add_paragraph()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    run = p.add_run()
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    p.add_run("目录")
    for i in range(10):
        if i == 2:
            para = doc.add_paragraph("第1章 监理工作内容")
            pPr = para._p.get_or_add_pPr()
            ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), "0"); pPr.append(ol)
        else:
            doc.add_paragraph("监理正文段落%d：全过程监理。" % i)
    if with_blank:
        doc.add_paragraph("")
        pb = doc.add_paragraph()
        br = OxmlElement("w:br"); br.set(qn("w:type"), "page")
        pb.add_run()._r.append(br)
        doc.add_paragraph("")
    doc.add_paragraph("本工程由总监理工程师赵监理负责，企业完成T18集成业绩项目。")  # 引用一致性明标名
    doc.add_paragraph("结尾段：本技术标制作完成。")
    doc.save(str(path))


def _remove_blank_pages(path):
    """修复：删除 双空夹分页 段落（空段+分页符空段+空段）。"""
    from docx import Document
    doc = Document(str(path))
    paras = doc.paragraphs
    targets = []
    n = len(paras)
    for i, p in enumerate(paras):
        xml = p._p.xml
        if 'w:br' in xml and 'w:type="page"' in xml:
            lo = max(0, i - 1)
            hi = min(n - 1, i + 1)
            if all(not (paras[j].text or "").strip() for j in (lo, hi)):
                targets += [lo, i, hi]
    for idx in sorted(set(targets), reverse=True):
        paras[idx]._p.getparent().remove(paras[idx]._p)
    doc.save(str(path))


class M10RecheckIntegrationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._run("init-enterprise", "--name", ENT)
        self.ent = self.root / ENT
        self.proj = self.ent / "项目级" / PROJECT
        self.tdir = self.proj / "招标解析"
        self.tdir.mkdir(parents=True)
        (self.tdir / ("投标要点_%s.json" % PROJECT)).write_text(
            json.dumps(_points(), ensure_ascii=False), encoding="utf-8")
        perf = self.proj / "业绩" / "T18集成业绩.pdf"
        perf.parent.mkdir(parents=True)
        perf.write_bytes(b"%PDF-1.4 fake")
        (self.tdir / ("素材清单_%s.json" % PROJECT)).write_text(
            json.dumps(_material(str(perf)), ensure_ascii=False), encoding="utf-8")
        tech = self.proj / "技术标"
        tech.mkdir(parents=True)
        (tech / "设置.json").write_text(
            json.dumps({"schema": "m6.settings/v1", "bidMode": "明标",
                        "pageConstraint": {"max": 60}}, ensure_ascii=False), encoding="utf-8")
        out = tech / "产出"
        out.mkdir(parents=True)
        self.docx = out / ("技术标_%s.docx" % PROJECT)
        _mk_tech_docx(self.docx, with_blank=True)

    def tearDown(self):
        self.tmp.cleanup()

    def _run(self, *argv):
        import io as _io
        from contextlib import redirect_stdout, redirect_stderr
        out, err, code = _io.StringIO(), _io.StringIO(), 0
        with redirect_stdout(out), redirect_stderr(err):
            try:
                bidcraft.main(["--root", str(self.root), "--enterprise", ENT] + list(argv))
            except SystemExit as e:
                code = e.code or 0
        return code, out.getvalue(), err.getvalue()

    def _check_path(self):
        return self.proj / "标书检查" / ("机器检查_%s_技术标.json" % PROJECT)

    def _blank_task_id(self):
        tp = self.proj / "模拟评标" / ("失分补救任务单_%s.json" % PROJECT)
        doc = json.loads(tp.read_text(encoding="utf-8"))
        for t in doc["tasks"]:
            if (t.get("检查快照") or {}).get("id") == "tech_blank_pages":
                return t["id"]
        raise AssertionError("任务单缺空白页检查任务")

    def test_recheck_passes_after_fix(self):
        # 1) 检查发现问题
        code, _, err = self._run("check-rule", "--project", PROJECT, "--out", str(self.docx),
                                 "--kind", "tech")
        self.assertEqual(code, 0, err)
        # 2) 生成任务单（check-only：无 M8 产物）
        code, _, err = self._run("remedy-gen", "--project", PROJECT, "--check", str(self._check_path()))
        self.assertEqual(code, 0, err)
        tid = self._blank_task_id()
        # 3) 勾选完成
        code, out, err = self._run("remedy-done", "--project", PROJECT, "--id", tid, "--note", "已删空段")
        self.assertEqual(code, 0, "%s | %s" % (out, err))
        self.assertIn("已勾选完成", out)
        # 4) 修复 docx → 重跑 check-rule
        _remove_blank_pages(self.docx)
        code, _, err = self._run("check-rule", "--project", PROJECT, "--out", str(self.docx),
                                 "--kind", "tech")
        self.assertEqual(code, 0, err)
        # 5) remedy-recheck：验证通过 + 报告产物
        code, out, err = self._run("remedy-recheck", "--project", PROJECT, "--check", str(self._check_path()))
        self.assertEqual(code, 0, "%s | %s" % (out, err))
        rp = self.proj / "模拟评标" / ("补救后重检报告_%s.json" % PROJECT)
        self.assertTrue(rp.is_file())
        r = json.loads(rp.read_text(encoding="utf-8"))
        self.assertIn("验证通过", r["summary"])
        hit = [x for x in r["results"] if x["id"] == tid]
        self.assertTrue(hit, "重检结果缺任务 %s" % tid)
        self.assertEqual(hit[0]["验证"], "通过")
        html = self.proj / "模拟评标" / ("补救后重检报告_%s.html" % PROJECT)
        self.assertTrue(html.is_file())
        self.assertIn("通过", html.read_text(encoding="utf-8"))
        # 6) 任务单写回：重检结果 + 已完成状态
        tp = self.proj / "模拟评标" / ("失分补救任务单_%s.json" % PROJECT)
        doc = json.loads(tp.read_text(encoding="utf-8"))
        t = next(x for x in doc["tasks"] if x["id"] == tid)
        self.assertEqual(t["重检结果"]["验证"], "通过")

    def test_recheck_fails_without_fix(self):
        code, _, err = self._run("check-rule", "--project", PROJECT, "--out", str(self.docx),
                                 "--kind", "tech")
        self.assertEqual(code, 0, err)
        code, _, err = self._run("remedy-gen", "--project", PROJECT, "--check", str(self._check_path()))
        self.assertEqual(code, 0, err)
        tid = self._blank_task_id()
        code, _, err = self._run("remedy-done", "--project", PROJECT, "--id", tid)
        self.assertEqual(code, 0, err)
        # 不修复 → 重检未通过
        code, _, err = self._run("remedy-recheck", "--project", PROJECT, "--check", str(self._check_path()))
        self.assertEqual(code, 0, err)
        rp = self.proj / "模拟评标" / ("补救后重检报告_%s.json" % PROJECT)
        r = json.loads(rp.read_text(encoding="utf-8"))
        hit = [x for x in r["results"] if x["id"] == tid]
        self.assertTrue(hit)
        self.assertEqual(hit[0]["验证"], "未通过")


if __name__ == "__main__":
    unittest.main()
