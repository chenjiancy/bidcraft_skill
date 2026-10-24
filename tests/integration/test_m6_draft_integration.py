# -*- coding: utf-8 -*-
"""M6 T14 集成测试（L2）：自动撰写引擎全链路 —— 机器草稿进精修链。

覆盖：
1) 全链路：tech-init → outline(--confirm) → ctx → tech-draft(--all) →
   tech-parse(逐章) → tech-rewrite(M6-REVISE 精修) → tech-check → tech-diff(裁决) →
   tech-export 产出可打开 docx。
2) 暗标链路：暗标模式下草稿无企业标识 → tech-parse 敏感词校验不拦截 → 可继续导出。
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(_HERE, "..", "..", "scripts"))

from docx import Document                     # noqa: E402
from _shared import core                      # noqa: E402
from m6_tech import store, cli as C           # noqa: E402

PROJECT = "T14集成项目"


def _points(bid_mode="明标"):
    rule = "技术标总页数不超过 60 页；%s" % (
        "技术标采用暗标，不得出现企业标识" if bid_mode == "暗标" else "按明标编制")
    return {
        "project": PROJECT,
        "modules": [
            {"id": "scoring", "title": "评分办法", "type": "mixed",
             "method": {"value": "综合评分法", "evidence": "e"},
             "score_items": [
                 {"item": "监理大纲（50分）", "weight": 50, "unit": "分", "evidence": "E1"},
                 {"item": "监理大纲·综合比较阐述（20分）", "weight": 20, "unit": "分", "evidence": "E2"}]},
            {"id": "reject", "title": "废标红线", "type": "items",
             "items": [{"type": "格式", "rule": rule, "evidence": "e"}]},
        ],
    }


def _draft():
    return {"chapters": [
        {"id": "ch-01", "no": "1", "title": "工程概况与监理范围", "level": 1,
         "scoreRefs": [{"idx": 0, "item": "监理大纲（50分）", "weight": 50}],
         "keyPoints": ["工程概况", "监理范围", "质量控制"], "wordCnt": 3000,
         "dependsOnMaterials": ["project_overview", "personnel", "performance", "enterprise_profile"]},
        {"id": "ch-02", "no": "2", "title": "质量管理与综合阐述", "level": 1,
         "scoreRefs": [{"idx": 1, "item": "监理大纲·综合比较阐述（20分）", "weight": 20}],
         "keyPoints": ["合同管理", "信息管理"], "wordCnt": 1000, "dependsOnMaterials": []},
    ]}


def _revised(md):
    """精修稿 = 草稿正文（剥掉 M6-COVERAGE）+ M6-REVISE 变更块（自述块唯一）。"""
    body = md.split("<!-- M6-COVERAGE")[0].rstrip()
    return body + '\n\n<!-- M6-REVISE\n{"changes":[{"where":"末尾","why":"agent 平台精修补强"}],"untouched":"其余段落未改动"}\n-->\n'


class DraftIntegrationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.ent, _ = core.Library(self.root).init_enterprise("T14集成企业")
        self.proj = self.ent / "项目级" / PROJECT
        self.tdir = self.proj / "招标解析"
        self.tdir.mkdir(parents=True)
        self._write_points()
        self._write_material()
        self.draft = self.proj / "章节树草案.json"
        self.draft.write_text(json.dumps(_draft(), ensure_ascii=False), encoding="utf-8")

    def tearDown(self):
        self.tmp.cleanup()

    def _write_points(self, bid_mode="明标"):
        (self.tdir / ("投标要点_%s.json" % PROJECT)).write_text(
            json.dumps(_points(bid_mode), ensure_ascii=False), encoding="utf-8")

    def _write_material(self):
        (self.tdir / ("素材清单_%s.json" % PROJECT)).write_text(
            json.dumps({"project": PROJECT,
                        "personnel": [{"name": "王五", "role": "总监理工程师"}],
                        "performance": [{"name": "T14集成业绩", "path": "业绩/x.pdf"}],
                        "企业基础信息": {"法定代表人姓名": "赵六"},
                        "qualification_required": [], "iso_certificates": [],
                        "honors": [], "items": []}, ensure_ascii=False), encoding="utf-8")

    def _tech(self):
        return store.tech_path(self.ent, PROJECT)

    def _args(self, **kw):
        base = {"root": str(self.root), "enterprise": "T14集成企业", "project": PROJECT}
        base.update(kw)
        from types import SimpleNamespace
        return SimpleNamespace(**base)

    def _prep(self, bid_mode="明标"):
        if bid_mode == "暗标":
            self._write_points("暗标")
        C.cmd_tech_init(self._args())
        C.cmd_tech_outline(self._args(file=str(self.draft), tdir="", confirm=True,
                                      expert_pack="", bid_mode=""))
        C.cmd_tech_ctx(self._args(chapter="", all=True))

    def _draft_all(self):
        from m6_tech import draft as D
        return D.cmd_tech_draft(self._args(chapter="", all=True, tdir="", json=False))

    def _parse_chain(self):
        """draft → parse（机器草稿直接落盘）→ rewrite（精修稿）→ check。"""
        from m6_tech import draft as D
        D.cmd_tech_draft(self._args(chapter="", all=True, tdir="", json=False))
        for cid in ("ch-01", "ch-02"):
            d = self._tech() / "章节" / ("%s.draft.md" % cid)
            C.cmd_tech_parse(self._args(chapter=cid, file=str(d), json=False))
            rev = self.proj / ("精修_%s.md" % cid)
            rev.write_text(_revised(d.read_text(encoding="utf-8")), encoding="utf-8")
            C.cmd_tech_rewrite(self._args(chapter=cid, file=str(rev)))
        C.cmd_tech_check(self._args(tdir="", chapter="", json=False))

    def _export(self, force=False):
        from m6_tech import export as E
        return E.cmd_tech_export(self._args(out="", override="", bid_mode="",
                                            force=force, template="", fill="", json=False))

    # -- 用例 -------------------------------------------------------------
    def test_full_chain_draft_to_export(self):
        """机器草稿 → parse/rewrite/check/diff/export 全链路闭环，docx 可打开。"""
        self._prep()
        self.assertEqual(self._draft_all(), 0)
        for cid in ("ch-01", "ch-02"):
            self.assertTrue((self._tech() / "章节" / ("%s.draft.md" % cid)).is_file(), cid)
        self._parse_chain()
        # diff 未裁决 → 拒绝导出
        from m6_tech import diff as D
        D.cmd_tech_diff(self._args(verdict="", confirm=False, json=False))
        with self.assertRaises(core.LibraryError):
            self._export()
        # 裁决 accept → 导出 docx
        jp = self._tech() / ("检查/双通道差异_%s.json" % PROJECT)
        n = len(json.loads(jp.read_text(encoding="utf-8"))["diffs"])
        vp = self.proj / "裁决.json"
        vp.write_text(json.dumps({"diffs": [{"idx": i, "verdict": "accept"}
                                            for i in range(n)]}, ensure_ascii=False),
                      encoding="utf-8")
        D.cmd_tech_diff(self._args(verdict=str(vp), confirm=True, json=False))
        self.assertEqual(self._export(), 0)
        out = self._tech() / "产出" / ("技术标_%s.docx" % PROJECT)
        self.assertTrue(out.is_file(), "导出 docx 缺失")
        doc = Document(str(out))
        self.assertGreater(len(doc.paragraphs), 0)
        self.assertIn("fldChar", doc.element.xml)          # 目录域

    def test_anbiao_chain_no_identity_leak(self):
        """暗标：草稿不含企业标识 → parse 敏感词校验不拦截 → 导出正文无企业名。"""
        self._prep("暗标")
        self.assertEqual(self._draft_all(), 0)
        md = (self._tech() / "章节" / "ch-01.draft.md").read_text(encoding="utf-8")
        for banned in ("T14集成企业", "王五", "赵六", "T14集成业绩"):
            self.assertNotIn(banned, md)
        self.assertIn("| 岗位 | 职责 |", md)
        self.assertIn("同类工程", md)
        self._parse_chain()                                  # 敏感词校验天然通过
        from m6_tech import diff as D
        D.cmd_tech_diff(self._args(verdict="", confirm=False, json=False))
        jp = self._tech() / ("检查/双通道差异_%s.json" % PROJECT)
        n = len(json.loads(jp.read_text(encoding="utf-8"))["diffs"])
        vp = self.proj / "裁决.json"
        vp.write_text(json.dumps({"diffs": [{"idx": i, "verdict": "accept"}
                                            for i in range(n)]}, ensure_ascii=False),
                      encoding="utf-8")
        D.cmd_tech_diff(self._args(verdict=str(vp), confirm=True, json=False))
        self.assertEqual(self._export(), 0)
        out = self._tech() / "产出" / ("技术标_%s.docx" % PROJECT)
        doc = Document(str(out))
        text = "\n".join(p.text for p in doc.paragraphs)
        for banned in ("T14集成企业", "王五", "赵六"):
            self.assertNotIn(banned, text)


if __name__ == "__main__":
    unittest.main()
