# -*- coding: utf-8 -*-
"""M9 T20 集成测试：真实企业两项目 —— 项目 A 走真实 check-rule 链（检查已过、评标未做→红灯），
项目 B 合成闭环（绿灯）→ m9-verify --all 企业汇总正确 + 企业汇总 HTML 下钻链接 + 单项目互链导航。"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(_HERE, "..", "..", "scripts"))

import bidcraft                       # noqa: E402
from docx.oxml import OxmlElement     # noqa: E402
from docx.oxml.ns import qn           # noqa: E402

ENT = "T20集成企业"
PROJ_A = "T20真实链项目"
PROJ_B = "T20闭环项目"


def _points(project):
    return {"project": project, "modules": [
        {"id": "scoring", "title": "评分办法", "type": "mixed",
         "method": {"value": "综合评分法", "evidence": "e"},
         "score_items": [{"item": "监理大纲（50分）", "weight": 50, "unit": "分", "evidence": "E1"}]},
        {"id": "reject", "title": "废标红线", "type": "items",
         "items": [{"type": "格式", "rule": "技术标总页数不超过 60 页；按明标编制", "evidence": "e"}]},
    ]}


def _material(project, perf_path):
    return {"project": project,
            "personnel": [{"name": "王监理", "role": "总监理工程师"}],
            "企业基础信息": {"法定代表人姓名": "刘总", "企业名称": ENT},
            "performance": [{"name": "%s业绩" % project, "path": perf_path}],
            "qualification_required": [], "iso_certificates": [], "honors": [], "items": []}


def _mk_tech_docx(path, project):
    from docx import Document
    doc = Document()
    doc.add_paragraph("封面：技术标")
    p = doc.add_paragraph()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    run = p.add_run()
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    p.add_run("目录")
    for i in range(8):
        if i == 2:
            para = doc.add_paragraph("第1章 监理工作内容")
            pPr = para._p.get_or_add_pPr()
            ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), "0"); pPr.append(ol)
        else:
            doc.add_paragraph("监理正文段落%d：全过程监理。" % i)
    doc.add_paragraph("本工程由总监理工程师王监理负责，企业完成%s业绩项目。" % project)
    doc.add_paragraph("结尾段：本技术标制作完成。")
    doc.save(str(path))


def _full_proj_artifacts(root, name):
    """合成闭环项目：全部产物 + M10 闭环。"""
    proj = root / ENT / "项目级" / name
    tdir = proj / "招标解析"; tdir.mkdir(parents=True)
    (tdir / ("投标要点_%s.json" % name)).write_text(json.dumps({"project": name}), encoding="utf-8")
    (tdir / ("素材清单_%s.json" % name)).write_text(json.dumps({"project": name}), encoding="utf-8")
    tech = proj / "技术标"; tech.mkdir(parents=True)
    (tech / "设置.json").write_text(json.dumps({"bidMode": "明标"}), encoding="utf-8")
    outd = tech / "产出"; outd.mkdir(parents=True)
    (outd / ("技术标_%s.docx" % name)).write_bytes(b"\x00docx")
    cdir = proj / "标书检查"; cdir.mkdir(parents=True)
    (cdir / ("机器检查_%s_技术标.json" % name)).write_text(json.dumps(
        {"kind": "tech", "检查": [{"id": "t", "ok": True, "issues": []}]}), encoding="utf-8")
    sdir = proj / "模拟评标"; sdir.mkdir(parents=True)
    (sdir / ("模拟评标_%s.json" % name)).write_text(json.dumps({"summary": {"issues": []}}), encoding="utf-8")
    (sdir / ("失分补救任务单_%s.json" % name)).write_text(json.dumps(
        {"project": name, "tasks": [{"id": "C01", "完成状态": "已完成"}]}), encoding="utf-8")
    (sdir / ("补救后重检报告_%s.json" % name)).write_text(json.dumps(
        {"project": name, "summary": {"已勾选任务": 1, "验证通过": 1}}), encoding="utf-8")
    return proj


class M9VerifyAllIntegrationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._run("init-enterprise", "--name", ENT)
        # 项目 A：真实链（M4 + 真实 docx + 真实 check-rule）
        proj = self.root / ENT / "项目级" / PROJ_A
        tdir = proj / "招标解析"; tdir.mkdir(parents=True)
        (tdir / ("投标要点_%s.json" % PROJ_A)).write_text(
            json.dumps(_points(PROJ_A), ensure_ascii=False), encoding="utf-8")
        perf = proj / "业绩" / ("%s业绩.pdf" % PROJ_A)
        perf.parent.mkdir(parents=True)
        perf.write_bytes(b"%PDF-1.4 fake")
        (tdir / ("素材清单_%s.json" % PROJ_A)).write_text(
            json.dumps(_material(PROJ_A, str(perf)), ensure_ascii=False), encoding="utf-8")
        tech = proj / "技术标"; tech.mkdir(parents=True)
        (tech / "设置.json").write_text(
            json.dumps({"bidMode": "明标", "pageConstraint": {"max": 60}}, ensure_ascii=False), encoding="utf-8")
        out = tech / "产出"; out.mkdir(parents=True)
        self.docx = out / ("技术标_%s.docx" % PROJ_A)
        _mk_tech_docx(self.docx, PROJ_A)
        # 项目 B：合成闭环
        _full_proj_artifacts(self.root, PROJ_B)

    def tearDown(self):
        self.tmp.cleanup()

    def _run(self, *argv):
        import io as _io
        from contextlib import redirect_stdout, redirect_stderr
        out, err, code = _io.StringIO(), _io.StringIO(), 0
        with redirect_stdout(out), redirect_stderr(err):
            try:
                bidcraft.main(["--root", str(self.root), "--enterprise", ENT] + list(argv))
            except SystemExit as e:
                code = e.code or 0
        return code, out.getvalue(), err.getvalue()

    def test_enterprise_all(self):
        # 项目 A 走真实 check-rule（机器检查产出）
        code, _, err = self._run("check-rule", "--project", PROJ_A, "--out", str(self.docx), "--kind", "tech")
        self.assertEqual(code, 0, err)

        # 企业级一键体检
        code, out, err = self._run("m9-verify", "--all")
        self.assertEqual(code, 0, "%s|%s" % (out, err))
        agg_path = self.root / ENT / "自检" / ("企业体检汇总_%s.json" % ENT)
        agg = json.loads(agg_path.read_text(encoding="utf-8"))
        s = agg["summary"]
        self.assertEqual(s["项目总数"], 2)
        by = {r["project"]: r for r in agg["projects"]}
        # B 合成闭环 → 绿灯
        self.assertEqual(by[PROJ_B]["总评"], "绿灯")
        # A 真实链：M7 已过但 M8/M10 未做 → 红灯
        self.assertEqual(by[PROJ_A]["总评"], "红灯")
        self.assertIn("m8_eval", by[PROJ_A]["未通过项"])
        self.assertEqual(s["含红灯项目"], [PROJ_A])

        # 三件套落盘
        for ext in ("json", "md", "html"):
            self.assertTrue((self.root / ENT / "自检" / ("企业体检汇总_%s.%s" % (ENT, ext))).is_file())
        # 企业汇总 HTML 下钻链接
        html = (self.root / ENT / "自检" / ("企业体检汇总_%s.html" % ENT)).read_text(encoding="utf-8")
        self.assertIn("../%s/自检/自检体检_%s.html" % (PROJ_B, PROJ_B), html)
        self.assertIn("../%s/自检/投标留痕包_%s.html" % (PROJ_A, PROJ_A), html)

        # 单项目互链导航（B 项目体检 + 留痕包）
        code, _, err = self._run("m9-verify", "--project", PROJ_B)
        self.assertEqual(code, 0, err)
        v = (self.root / ENT / "项目级" / PROJ_B / "自检" / ("自检体检_%s.html" % PROJ_B)).read_text(encoding="utf-8")
        self.assertIn("互链", v)
        self.assertIn("投标留痕包_%s.html" % PROJ_B, v)
        self.assertIn("企业体检汇总_%s.html" % ENT, v)
        self.assertIn("补救后重检报告_%s.html" % PROJ_B, v)


if __name__ == "__main__":
    unittest.main()
