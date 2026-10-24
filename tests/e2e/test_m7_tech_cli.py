# -*- coding: utf-8 -*-
"""M7 T15 端到端测试（L3）：技术标检查 CLI 全链路 ——
check-rule(--kind tech) → check-parse(通道B) → check-diff → check-report，tag=技术标。
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(_HERE, "..", "..", "scripts"))

import bidcraft                        # noqa: E402
from _shared import core               # noqa: E402

PROJECT = "T15端到端项目"
ENT = "T15端到端企业"


def _points():
    return {"project": PROJECT, "modules": [
        {"id": "scoring", "title": "评分办法", "type": "mixed",
         "method": {"value": "综合评分法", "evidence": "e"},
         "score_items": [{"item": "监理大纲（50分）", "weight": 50, "unit": "分", "evidence": "E1"}]},
        {"id": "reject", "title": "废标红线", "type": "items",
         "items": [{"type": "格式", "rule": "技术标总页数不超过 20 页；按明标编制", "evidence": "e"}]},
    ]}


def _material(perf_path):
    return {"project": PROJECT,
            "personnel": [{"name": "周十一", "role": "总监理工程师"}],
            "企业基础信息": {"法定代表人姓名": "吴十二", "企业名称": ENT},
            "performance": [{"name": "端到端业绩", "path": perf_path}],
            "qualification_required": [], "iso_certificates": [], "honors": [], "items": []}


def _mk_tech_docx(path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    doc.add_paragraph("封面：技术标")
    p = doc.add_paragraph()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    run = p.add_run()
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    p.add_run("目录")
    for i in range(20):
        if i == 2:
            para = doc.add_paragraph("第1章 监理工作内容")
            pPr = para._p.get_or_add_pPr()
            ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), "0"); pPr.append(ol)
        else:
            doc.add_paragraph("监理正文段落%d：全过程监理。" % i)
    doc.add_paragraph("结尾段：本技术标制作完成。")
    doc.save(str(path))


def _agent_points_json():
    return {
        "project": PROJECT, "channel": "B", "审查日期": "2026-10-24",
        "结论": "通过", "废标红线": [],
        "维度明细": [{"维度": "排版", "风险等级": "🟢", "问题描述": "结构完整", "来源追溯": "检查", "整改建议": "无"}],
        "整改清单": [], "通道A差异说明": "与机器检查一致", "总结": "技术标结构完整",
    }


class M7TechE2ETest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._run("init-enterprise", "--name", ENT)
        self.ent = self.root / ENT
        self.proj = self.ent / "项目级" / PROJECT
        self.tdir = self.proj / "招标解析"
        self.tdir.mkdir(parents=True)
        (self.tdir / ("投标要点_%s.json" % PROJECT)).write_text(
            json.dumps(_points(), ensure_ascii=False), encoding="utf-8")
        perf = self.proj / "业绩" / "端到端业绩.pdf"
        perf.parent.mkdir(parents=True)
        perf.write_bytes(b"%PDF-1.4 fake")
        (self.tdir / ("素材清单_%s.json" % PROJECT)).write_text(
            json.dumps(_material(str(perf)), ensure_ascii=False), encoding="utf-8")
        tech = self.proj / "技术标"
        tech.mkdir(parents=True)
        (tech / "设置.json").write_text(
            json.dumps({"schema": "m6.settings/v1", "bidMode": "明标",
                        "pageConstraint": {"max": 20}}, ensure_ascii=False), encoding="utf-8")
        out = tech / "产出"
        out.mkdir(parents=True)
        self.docx = out / ("技术标_%s.docx" % PROJECT)
        _mk_tech_docx(self.docx)

    def tearDown(self):
        self.tmp.cleanup()

    def _run(self, *argv):
        import io as _io
        from contextlib import redirect_stdout, redirect_stderr
        out, err, code = _io.StringIO(), _io.StringIO(), 0
        with redirect_stdout(out), redirect_stderr(err):
            try:
                bidcraft.main(["--root", str(self.root), "--enterprise", ENT] + list(argv))
            except SystemExit as e:
                code = e.code or 0
        return code, out.getvalue(), err.getvalue()

    def test_tech_check_cli_full_chain(self):
        # 1) check-rule --kind tech（auto 推断也覆盖）
        code, out, _ = self._run("check-rule", "--project", PROJECT, "--out", str(self.docx),
                                 "--kind", "tech")
        self.assertEqual(code, 0, out)
        self.assertIn("技术标", out)
        jp = self.proj / "标书检查" / ("机器检查_%s_技术标.json" % PROJECT)
        self.assertTrue(jp.is_file())
        doc = json.loads(jp.read_text(encoding="utf-8"))
        self.assertEqual(doc["kind"], "tech")
        ids = [c["id"] for c in doc["检查"]]
        for want in ("tech_structure", "anbiao_leak", "tech_pages", "material"):
            self.assertIn(want, ids)
        self.assertNotIn("quote", ids)

        # 2) check-parse 通道B（agent 要点）→ tag 一致
        md_src = self.proj / "agent_要点.md"
        md_src.write_text("# M7 通道B 检查要点\n结论：通过\n", encoding="utf-8")
        js_src = self.proj / "agent_要点.json"
        js_src.write_text(json.dumps(_agent_points_json(), ensure_ascii=False), encoding="utf-8")
        code, _, err = self._run("check-parse", "--project", PROJECT,
                                 "--md", str(md_src), "--json-file", str(js_src),
                                 "--tag", "技术标")
        self.assertEqual(code, 0, err)

        # 3) check-diff（双通道，tag 对齐）
        code, _, err = self._run("check-diff", "--project", PROJECT, "--tag", "技术标")
        self.assertEqual(code, 0, err)
        dj = self.proj / "标书检查" / ("双通道检查差异_%s_技术标.json" % PROJECT)
        self.assertTrue(dj.is_file())

        # 4) check-report HTML
        code, _, err = self._run("check-report", "--project", PROJECT, "--tag", "技术标")
        self.assertEqual(code, 0, err)
        rp = self.proj / "标书检查" / ("标书检查报告_%s_技术标.html" % PROJECT)
        self.assertTrue(rp.is_file())
        html = rp.read_text(encoding="utf-8")
        self.assertIn("技术标", html)
        self.assertIn("tech_structure", html)

    def test_tech_check_auto_kind_by_path(self):
        code, out, _ = self._run("check-rule", "--project", PROJECT, "--out", str(self.docx))
        self.assertEqual(code, 0, out)
        self.assertIn("技术标", out)          # auto 按路径推断 → tag 默认「技术标」
        jp = self.proj / "标书检查" / ("机器检查_%s_技术标.json" % PROJECT)
        self.assertTrue(jp.is_file())


if __name__ == "__main__":
    unittest.main()
