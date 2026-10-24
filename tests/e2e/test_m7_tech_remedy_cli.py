# -*- coding: utf-8 -*-
"""M7 T17 端到端测试（L3）：技术标检查扩充 + M10 补救任务单联动 ——
check-rule --kind tech（含新 3 项：图表位置/空白页/引用一致性）落盘机器检查
→ remedy-gen --check <机器检查json> 生成含「技术标检查整改」任务的任务单（与失分补救并列）。
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(_HERE, "..", "..", "scripts"))

import bidcraft                        # noqa: E402
from docx.oxml import OxmlElement      # noqa: E402
from docx.oxml.ns import qn            # noqa: E402

PROJECT = "T17端到端项目"
ENT = "T17端到端企业"


def _points():
    return {"project": PROJECT, "modules": [
        {"id": "scoring", "title": "评分办法", "type": "mixed",
         "method": {"value": "综合评分法", "evidence": "e"},
         "score_items": [{"item": "监理大纲（50分）", "weight": 50, "unit": "分", "evidence": "E1"}]},
        {"id": "reject", "title": "废标红线", "type": "items",
         "items": [{"type": "格式", "rule": "技术标总页数不超过 60 页；按明标编制", "evidence": "e"}]},
    ]}


def _material(perf_path):
    return {"project": PROJECT,
            "personnel": [{"name": "钱十三", "role": "总监理工程师"}],
            "企业基础信息": {"法定代表人姓名": "孙十四", "企业名称": ENT},
            "performance": [{"name": "T17端到端业绩", "path": perf_path}],
            "qualification_required": [], "iso_certificates": [], "honors": [], "items": []}


def _mk_tech_docx(path):
    """合成技术标 docx：TOC 域 + 大纲 + 正文，含（a）正文中段插图（图表位置违规）、
    （b）双空夹分页（空白页）、（c）正文不出现清单业绩名/人员名（引用一致性缺失）。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph("封面：技术标")
    p = doc.add_paragraph()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    run = p.add_run()
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    p.add_run("目录")
    for i in range(12):
        if i == 2:
            para = doc.add_paragraph("第1章 监理工作内容")
            pPr = para._p.get_or_add_pPr()
            ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), "0"); pPr.append(ol)
        else:
            doc.add_paragraph("监理正文段落%d：全过程监理。" % i)
    # (a) 正文中段插图
    img = doc.add_paragraph("图 1：工程平面示意")
    run = img.add_run()
    drawing = OxmlElement("w:drawing")
    inline = OxmlElement("wp:inline")
    inline.append(OxmlElement("wp:extent"))
    drawing.append(inline)
    run._r.append(drawing)
    doc.add_paragraph("监理正文段落-图后：继续正文。")
    # (b) 双空夹分页（空白页）
    doc.add_paragraph("")
    pb = doc.add_paragraph()
    br = OxmlElement("w:br"); br.set(qn("w:type"), "page")
    pb.add_run()._r.append(br)
    doc.add_paragraph("")
    doc.add_paragraph("结尾段：本技术标制作完成（正文未引用清单业绩与人员）。")
    doc.save(str(path))


class M7TechRemedyE2ETest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self._run("init-enterprise", "--name", ENT)
        self.ent = self.root / ENT
        self.proj = self.ent / "项目级" / PROJECT
        self.tdir = self.proj / "招标解析"
        self.tdir.mkdir(parents=True)
        (self.tdir / ("投标要点_%s.json" % PROJECT)).write_text(
            json.dumps(_points(), ensure_ascii=False), encoding="utf-8")
        perf = self.proj / "业绩" / "T17端到端业绩.pdf"
        perf.parent.mkdir(parents=True)
        perf.write_bytes(b"%PDF-1.4 fake")
        (self.tdir / ("素材清单_%s.json" % PROJECT)).write_text(
            json.dumps(_material(str(perf)), ensure_ascii=False), encoding="utf-8")
        tech = self.proj / "技术标"
        tech.mkdir(parents=True)
        (tech / "设置.json").write_text(
            json.dumps({"schema": "m6.settings/v1", "bidMode": "明标",
                        "pageConstraint": {"max": 60}}, ensure_ascii=False), encoding="utf-8")
        out = tech / "产出"
        out.mkdir(parents=True)
        self.docx = out / ("技术标_%s.docx" % PROJECT)
        _mk_tech_docx(self.docx)

    def tearDown(self):
        self.tmp.cleanup()

    def _run(self, *argv):
        import io as _io
        from contextlib import redirect_stdout, redirect_stderr
        out, err, code = _io.StringIO(), _io.StringIO(), 0
        with redirect_stdout(out), redirect_stderr(err):
            try:
                bidcraft.main(["--root", str(self.root), "--enterprise", ENT] + list(argv))
            except SystemExit as e:
                code = e.code or 0
        return code, out.getvalue(), err.getvalue()

    def test_check_rule_finds_t17_issues_then_remedy_link(self):
        # 1) check-rule --kind tech：三个新检查项命中
        code, out, _ = self._run("check-rule", "--project", PROJECT, "--out", str(self.docx),
                                 "--kind", "tech")
        self.assertEqual(code, 0, out)
        jp = self.proj / "标书检查" / ("机器检查_%s_技术标.json" % PROJECT)
        self.assertTrue(jp.is_file())
        doc = json.loads(jp.read_text(encoding="utf-8"))
        ids = {c["id"]: c for c in doc["检查"]}
        for want in ("tech_figure_placement", "tech_blank_pages", "tech_ref_consistency"):
            self.assertIn(want, ids)
        self.assertFalse(ids["tech_figure_placement"]["ok"],
                         ids["tech_figure_placement"]["issues"])
        self.assertFalse(ids["tech_blank_pages"]["ok"], ids["tech_blank_pages"]["issues"])
        self.assertFalse(ids["tech_ref_consistency"]["ok"], ids["tech_ref_consistency"]["issues"])

        # 2) remedy-gen --check 机器检查 json：检查整改任务并入任务单
        code, out, err = self._run("remedy-gen", "--project", PROJECT, "--check", str(jp))
        self.assertEqual(code, 0, "%s | %s" % (out, err))
        tp = self.proj / "模拟评标" / ("失分补救任务单_%s.json" % PROJECT)
        self.assertTrue(tp.is_file(), "任务单 json 缺失（remedy-gen 依赖 M8 产物，应自动创建）")
        doc = json.loads(tp.read_text(encoding="utf-8"))
        tasks = doc["tasks"]
        check_tasks = [t for t in tasks if t["来源"].startswith("M7-")]
        self.assertGreaterEqual(len(check_tasks), 3, "检查整改任务应覆盖 3 个未通过检查项")
        self.assertTrue(all(t["类型"] == "格式合规" for t in check_tasks))
        self.assertTrue(all(t["动作"] == "技术标检查整改" for t in check_tasks))
        self.assertIn("check-rule", check_tasks[0]["验证方式"])
        self.assertIn("生成依据", doc)
        self.assertIn("M7 机器检查", doc["生成依据"])


if __name__ == "__main__":
    unittest.main()
