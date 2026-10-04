# -*- coding: utf-8 -*-
"""docx_util 公共工具单测：跨 run 替换 / 整段写入 / 空段落 / 单元格。"""
import os
import sys
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))

from docx import Document  # noqa: E402

from _shared import docx_util  # noqa: E402


def _split_run(p, cuts):
    """把段落文本按字符下标切成多个 run（模拟 Word 跨 run 拆分）。"""
    full = p.text
    runs = p.runs
    if not runs:
        return
    parts = []
    prev = 0
    for c in cuts:
        parts.append(full[prev:c])
        prev = c
    parts.append(full[prev:])
    for r, part in zip(runs, parts):
        r.text = part
    if len(parts) > len(runs):
        for part in parts[len(runs):]:
            p.add_run(part)


class TestParaText(unittest.TestCase):
    def test_cross_run_concat(self):
        doc = Document()
        p = doc.add_paragraph("投标人名称：【企业名称】")
        _split_run(p, [3, 5, 9])
        self.assertEqual(docx_util.para_text(p), "投标人名称：【企业名称】")


class TestSetParaText(unittest.TestCase):
    def test_keeps_first_run_rpr(self):
        doc = Document()
        p = doc.add_paragraph("old text")
        p.runs[0].bold = True
        docx_util.set_para_text(p, "新文本")
        self.assertEqual(p.text, "新文本")
        self.assertTrue(p.runs[0].bold)
        self.assertEqual(len(p.runs), 1)

    def test_empty_paragraph_creates_run(self):
        doc = Document()
        p = doc.add_paragraph()
        docx_util.set_para_text(p, "写入")
        self.assertEqual(p.text, "写入")

    def test_multiple_runs_collapsed(self):
        doc = Document()
        p = doc.add_paragraph("甲")
        p.add_run("乙")
        p.add_run("丙")
        docx_util.set_para_text(p, "一")
        self.assertEqual(p.text, "一")
        self.assertEqual(len(p.runs), 1)


class TestReplaceInPara(unittest.TestCase):
    def test_cross_run_replace(self):
        doc = Document()
        p = doc.add_paragraph("致招标人：投标人（盖章）")
        _split_run(p, [4, 6])          # 占位词跨 run
        n = docx_util.replace_in_para(p, "投标人", "【企业名称】")
        self.assertEqual(n, 1)
        self.assertEqual(p.text, "致招标人：【企业名称】（盖章）")

    def test_replace_all_occurrences(self):
        doc = Document()
        p = doc.add_paragraph("甲方【日期】与【日期】乙方")
        _split_run(p, [4, 7])
        n = docx_util.replace_in_para(p, "【日期】", "2026年10月9日")
        self.assertEqual(n, 2)
        self.assertEqual(p.text, "甲方2026年10月9日与2026年10月9日乙方")

    def test_no_match_returns_zero_unchanged(self):
        doc = Document()
        p = doc.add_paragraph("原文不动")
        self.assertEqual(docx_util.replace_in_para(p, "不存在", "x"), 0)
        self.assertEqual(p.text, "原文不动")

    def test_empty_old_returns_zero(self):
        doc = Document()
        p = doc.add_paragraph("abc")
        self.assertEqual(docx_util.replace_in_para(p, "", "x"), 0)
        self.assertEqual(p.text, "abc")


class TestSetCellText(unittest.TestCase):
    def test_cell_clear_write(self):
        doc = Document()
        tbl = doc.add_table(1, 1)
        cell = tbl.rows[0].cells[0]
        cell.text = "旧值"
        docx_util.set_cell_text(cell, "新值")
        self.assertEqual(cell.text, "新值")
        self.assertEqual(len(cell.paragraphs[0].runs), 1)


if __name__ == "__main__":
    unittest.main()
