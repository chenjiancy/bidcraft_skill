# -*- coding: utf-8 -*-
"""L3：gen_toc 单测（TOC 域 / Heading 样式）。"""
import os
import sys
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))

from docx import Document

from m5_project import gen_toc


class TestTocField(unittest.TestCase):
    def test_add_toc_field(self):
        doc = Document()
        doc.add_paragraph("正文占位")
        r = gen_toc.add_toc_field(doc, title="目  录")
        self.assertTrue(r["标题"])
        self.assertTrue(r["域"])
        xml = doc.element.body.xml
        # 标题段存在
        self.assertIn("目  录", xml)
        # TOC 域指令完整
        self.assertIn("fldCharType=\"begin\"", xml)
        self.assertIn('TOC \\o "1-3"', xml)          # instrText 中为单个反斜杠
        self.assertIn("fldCharType=\"separate\"", xml)
        self.assertIn("fldCharType=\"end\"", xml)
        self.assertIn("w:dirty=\"true\"", xml)

    def test_toc_field_at_doc_start(self):
        """目录标题 + TOC 域应位于文档最前（合稿整本第一页）。"""
        doc = Document()
        doc.add_paragraph("A")
        doc.add_paragraph("B")
        gen_toc.add_toc_field(doc)
        first_paras = list(doc.element.body.iterchildren())[:4]
        texts = []
        for el in first_paras:
            if el.tag == qn_para():
                texts.append("".join(t.text or "" for t in el.iter() if t.tag == qn_t()))
        self.assertEqual(texts[0], "目  录")
        self.assertIn("（打开文档后按 F9 更新目录）", " ".join(texts[1:]))

    def test_add_heading_style(self):
        doc = Document()
        # python-docx 默认模板已含 Heading 1-9 样式 → 返回 True（已存在），且带 outlineLvl 可被 TOC 索引
        self.assertTrue(gen_toc.add_heading_style(doc, 1))
        st = doc.styles["Heading 1"]
        self.assertIn("outlineLvl", st.element.get_or_add_pPr().xml)


def qn_para():
    from docx.oxml.ns import qn as _qn
    return _qn("w:p")


def qn_t():
    from docx.oxml.ns import qn as _qn
    return _qn("w:t")


if __name__ == "__main__":
    unittest.main()
