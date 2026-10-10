# -*- coding: utf-8 -*-
"""page_fit 排版收敛单元测试（_apply_round 不依赖 Word COM，可单测）。

收敛语义（用户 2026-10-10 澄清）：文字可删空行/调行距；表格可删空白行/调行高；
**不改变字号、不改变页边距**（保持企业模板格式契约值）；不得影响阅读与美观。
"""
import os
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Cm, Pt

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "scripts"))
from m5_project import page_fit  # noqa: E402


class ApplyRoundTest(unittest.TestCase):
    def _mk_doc(self):
        doc = Document()
        p = doc.add_paragraph()
        r = p.add_run("测试文字")
        r.font.size = Pt(12)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(18)
        doc.add_paragraph()                       # 空段落（应被删除）
        doc.add_paragraph()                       # 空段落（应被删除）
        t = doc.add_table(rows=3, cols=2)
        t.cell(0, 0).text = "表头A"
        t.cell(0, 1).text = "表头B"
        t.cell(1, 0).text = "数据1"
        t.cell(1, 1).text = "数据2"
        # 第 2 行为空白行（应被删除）
        t.cell(2, 0).text = ""
        t.cell(2, 1).text = ""
        return doc

    def test_blank_paras_removed(self):
        doc = self._mk_doc()
        changed = page_fit._apply_round(doc, 0)
        self.assertTrue(changed)
        # 空段落被删除：剩余正文段数 = 1（原 3 段 → 删 2）
        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(len(texts), 1)
        self.assertEqual(texts[0], "测试文字")

    def test_line_spacing_and_space_shrunk(self):
        doc = self._mk_doc()
        page_fit._apply_round(doc, 0)
        p = doc.paragraphs[0]
        # V5：round0 行距=固定 28pt（EMU 355600），段距≤8pt
        self.assertEqual(p.paragraph_format.line_spacing_rule, WD_LINE_SPACING.EXACTLY)
        self.assertLessEqual(getattr(p.paragraph_format.line_spacing, "pt",
                                     p.paragraph_format.line_spacing or 0), 28.0)
        self.assertLessEqual(p.paragraph_format.space_after.pt, 8.0)

    def test_blank_table_row_removed(self):
        doc = self._mk_doc()
        page_fit._apply_round(doc, 0)
        t = doc.tables[0]
        rows_text = [tuple(c.text for c in row.cells) for row in t.rows]
        self.assertNotIn(("", ""), rows_text)     # 空白行已删除
        self.assertEqual(len(t.rows), 2)

    def test_font_size_unchanged(self):
        doc = self._mk_doc()
        before = doc.paragraphs[0].runs[0].font.size
        page_fit._apply_round(doc, 5)             # 多轮收敛
        after = doc.paragraphs[0].runs[0].font.size
        self.assertEqual(before, after)           # 字号不变

    def test_margin_unchanged(self):
        doc = self._mk_doc()
        sec = doc.sections[0]
        before = (sec.top_margin, sec.left_margin)
        page_fit._apply_round(doc, 6)
        sec = doc.sections[0]
        self.assertEqual(before, (sec.top_margin, sec.left_margin))  # 页边距不变

    def test_row_height_shrunk(self):
        doc = Document()
        t = doc.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "x"
        row = t.rows[0]
        # 显式设置超大行高（twips；1cm ≈ 567 twips → 3cm ≈ 1701）
        from docx.oxml.ns import qn
        trPr = row._tr.get_or_add_trPr()
        th = trPr.makeelement(qn("w:trHeight"), {qn("w:val"): "1701", qn("w:hRule"): "atLeast"})
        trPr.append(th)
        page_fit._apply_round(doc, 0)             # round0 行高上限 1.5cm ≈ 850 twips
        trPr = row._tr.trPr
        v = int(trPr.find(qn("w:trHeight")).get(qn("w:val")))
        self.assertLessEqual(v, 851)

    def test_shrink_to_fit_rounds_save(self):
        # 收敛路径不查页数（无 Word），验证保存与返回结构
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "t.docx"
            doc = self._mk_doc()
            doc.save(path)
            from unittest.mock import patch
            with patch.object(page_fit, "page_count", return_value=1):
                res = page_fit.shrink_to_fit(path)
            self.assertEqual(res["达成"], True)
            self.assertEqual(res["页数"], 1)
            re_doc = Document(str(path))
            self.assertTrue(len(re_doc.paragraphs) >= 1)


if __name__ == "__main__":
    unittest.main()
