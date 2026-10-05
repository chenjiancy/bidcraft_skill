# -*- coding: utf-8 -*-
"""unit · gen_diff 内容校验差异清单（D16–D20）。"""
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))

from docx import Document  # noqa: E402

from m5_project import gen_diff as gd  # noqa: E402
from m5_project.gen_blocks import _extract_blocks  # noqa: E402


def _mk_doc(paras, tables=None):
    """构造 docx：paras=[文本...]；tables=[[[行文本...]]] 可选。"""
    doc = Document()
    for t in paras:
        doc.add_paragraph(t)
    for tbl in tables or []:
        tab = doc.add_table(rows=len(tbl), cols=len(tbl[0]))
        for ri, row in enumerate(tbl):
            for ci, val in enumerate(row):
                tab.rows[ri].cells[ci].text = val
    return doc


class TestDiffTemplateVsContract:
    def setup_method(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def teardown_method(self):
        self.tmp.cleanup()

    def test_equal_no_diff(self):
        """模板与契约一致 → 无内容差异。"""
        tpl = self.root / "tpl.docx"
        _mk_doc(["单位名称：示例建设工程监理有限公司",
                 "地址：示例市"]).save(str(tpl))
        src = self.root / "src.docx"
        _mk_doc(["单位名称：示例建设工程监理有限公司",
                 "地址：示例市"]).save(str(src))
        blocks, _ = _extract_blocks(src)
        diffs = gd.diff_template_vs_contract(tpl, blocks, [0, 1], "F01", "封面.docx")
        assert diffs == []

    def test_content_diff_detected(self):
        """模板与契约实质文本差异 → content_diff，建议 update。"""
        tpl = self.root / "tpl.docx"
        _mk_doc(["单位名称：示例建设工程监理有限公司"]).save(str(tpl))
        src = self.root / "src.docx"
        _mk_doc(["单位名称：示例某某公司"]).save(str(src))
        blocks, _ = _extract_blocks(src)
        diffs = gd.diff_template_vs_contract(tpl, blocks, [0, 0], "F01", "封面.docx")
        cd = [d for d in diffs if d["类别"] == "content_diff"]
        assert len(cd) == 1
        assert cd[0]["建议"] == "update"

    def test_placeholder_zone_skipped(self):
        """模板含【占位】或契约为下划线待填区 → placeholder_zone，不属内容差异。"""
        tpl = self.root / "tpl.docx"
        _mk_doc(["单位名称：示例建设工程监理有限公司"]).save(str(tpl))
        src = self.root / "src.docx"
        _mk_doc(["单位名称：____________"]).save(str(src))
        blocks, _ = _extract_blocks(src)
        diffs = gd.diff_template_vs_contract(tpl, blocks, [0, 0], "F01", "封面.docx")
        assert diffs and all(d["类别"] == "placeholder_zone" for d in diffs)

        tpl2 = self.root / "tpl2.docx"
        _mk_doc(["项目名称：【项目名称】"]).save(str(tpl2))
        src2 = self.root / "src2.docx"
        _mk_doc(["项目名称：示例化工园尾水水质提升工程"]).save(str(src2))
        blocks2, _ = _extract_blocks(src2)
        diffs2 = gd.diff_template_vs_contract(tpl2, blocks2, [0, 0], "F01", "封面.docx")
        assert all(d["类别"] == "placeholder_zone" for d in diffs2)

    def test_table_cell_diff(self):
        """表格单元格差异 → content_diff；占位/待填格跳过。"""
        tpl = self.root / "tpl.docx"
        _mk_doc([], tables=[[["姓名", "性别"], ["张三", "男"]]]).save(str(tpl))
        src = self.root / "src.docx"
        _mk_doc([], tables=[[["姓名", "性别"], ["李四", "女"]]]).save(str(src))
        blocks, _ = _extract_blocks(src)
        diffs = gd.diff_template_vs_contract(tpl, blocks, [0, 0], "F06", "资格证明.docx")
        cd = [d for d in diffs if d["类别"] == "content_diff"]
        assert len(cd) == 2                      # 张三/李四、男/女

    def test_manifest_and_decisions(self):
        """清单落盘 + 用户改决策后 load_decisions 正确分组。"""
        tpl = self.root / "tpl.docx"
        _mk_doc(["单位名称：示例建设工程监理有限公司", "固定条款：A"]).save(str(tpl))
        src = self.root / "src.docx"
        _mk_doc(["单位名称：示例某某公司", "固定条款：A"]).save(str(src))
        blocks, _ = _extract_blocks(src)
        diffs = gd.diff_template_vs_contract(tpl, blocks, [0, 1], "F01", "封面.docx")
        md, js = gd.write_diff_manifest(diffs, "测试项目", self.root)
        assert md.is_file() and js.is_file()
        keep, update = gd.load_decisions(js)
        assert len(update) == 1 and keep == []

    def test_apply_keep_writes_back(self):
        """用户决策 keep → 生成稿写回模板文本。"""
        tpl = self.root / "tpl.docx"
        _mk_doc(["单位名称：示例建设工程监理有限公司"]).save(str(tpl))
        src = self.root / "src.docx"
        _mk_doc(["单位名称：示例某某公司"]).save(str(src))
        blocks, _ = _extract_blocks(src)
        diffs = gd.diff_template_vs_contract(tpl, blocks, [0, 0], "F01", "封面.docx")
        # 生成稿 = 契约文字
        gen = self.root / "gen.docx"
        _mk_doc(["单位名称：示例某某公司"]).save(str(gen))
        # 用户把该项决策改为 keep
        diffs[0]["建议"] = "keep"
        js = self.root / "diff.json"
        js.write_text(__import__("json").dumps({"差异": diffs}, ensure_ascii=False),
                      encoding="utf-8")
        keep, _ = gd.load_decisions(js)
        n, missing = gd.apply_keep_to_docx(gen, keep)
        assert n == 1 and missing == []
        out = Document(str(gen))
        assert "示例建设工程监理有限公司" in out.paragraphs[0].text
