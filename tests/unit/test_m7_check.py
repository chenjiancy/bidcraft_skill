# -*- coding: utf-8 -*-
"""M7 交标前质量检查单测：检查表结构 / 三道检查逻辑（合成 docx 与 zip）。"""
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))

from m7_check import checker as m7  # noqa: E402


class TestChecksTable(unittest.TestCase):
    def test_table_rows_have_schema(self):
        self.assertGreaterEqual(len(m7.CHECKS), 6)
        for item in m7.CHECKS:
            self.assertIn(item["id"], ("format", "placeholder", "sensitive", "word_open",
                                       "pages", "frozen"))
            self.assertTrue(callable(getattr(m7, item["fn"], None)),
                            "检查实现缺失: %s" % item["fn"])
            self.assertIn("level", item)

    def test_pages_skips_without_com(self):
        """⑫ 无 pywin32 环境页数检查跳过且不降级 ok。"""
        old = m7.HAVE_COM
        try:
            m7.HAVE_COM = False
            with tempfile.TemporaryDirectory() as td:
                out, tpl = Path(td) / "out", Path(td) / "tpl"
                out.mkdir(); tpl.mkdir()
                ok, issues = m7.check_pages(out, tpl)
                self.assertTrue(ok)
                self.assertTrue(issues and "跳过" in issues[0])
        finally:
            m7.HAVE_COM = old

    def test_word_open_skips_without_com(self):
        """⑫ 无 pywin32 环境跳过且不降级 ok。"""
        old = m7.HAVE_COM
        try:
            m7.HAVE_COM = False
            with tempfile.TemporaryDirectory() as td:
                out, tpl = Path(td) / "out", Path(td) / "tpl"
                out.mkdir(); tpl.mkdir()
                (out / "a.docx").write_bytes(b"x")
                ok, issues = m7.check_word_open(out, tpl)
                self.assertTrue(ok)
                self.assertTrue(any("跳过" in x for x in issues))
        finally:
            m7.HAVE_COM = old


class TestCheckPlaceholder(unittest.TestCase):
    def _mk_doc(self, path, paras, cells=None):
        from docx import Document
        doc = Document()
        for t in paras:
            doc.add_paragraph(t)
        if cells:
            tbl = doc.add_table(1, 1)
            tbl.rows[0].cells[0].text = cells
        doc.save(str(path))

    def test_image_placeholder_flagged(self):
        with tempfile.TemporaryDirectory() as td:
            out, tpl = Path(td) / "out", Path(td) / "tpl"
            out.mkdir(); tpl.mkdir()
            self._mk_doc(out / "a.docx", ["正文", "【图片：营业执照】"])
            ok, issues = m7.check_placeholder(out, tpl)
            self.assertFalse(ok)
            self.assertTrue(any("图片占位未填充" in x for x in issues))

    def test_allowlisted_pending_not_flagged(self):
        with tempfile.TemporaryDirectory() as td:
            out, tpl = Path(td) / "out", Path(td) / "tpl"
            out.mkdir(); tpl.mkdir()
            self._mk_doc(out / "a.docx", ["正文", "【投标总价（元）】"])
            (out / "待补字段清单.md").write_text(
                "- 【投标总价（元）】：报价策略待定（保留占位待补）", encoding="utf-8")
            ok, issues = m7.check_placeholder(out, tpl)
            self.assertTrue(ok, issues)

    def test_outside_placeholder_flagged(self):
        with tempfile.TemporaryDirectory() as td:
            out, tpl = Path(td) / "out", Path(td) / "tpl"
            out.mkdir(); tpl.mkdir()
            self._mk_doc(out / "a.docx", ["【未登记占位】"])
            (out / "待补字段清单.md").write_text("- 【已登记】", encoding="utf-8")
            ok, issues = m7.check_placeholder(out, tpl)
            self.assertFalse(ok)
            self.assertTrue(any("清单外残留占位" in x for x in issues))


class TestCheckFormat(unittest.TestCase):
    def _mk_doc(self, path, paras=("正文",), tables=0):
        from docx import Document
        doc = Document()
        for t in paras:
            doc.add_paragraph(t)
        for _ in range(tables):
            doc.add_table(1, 1)
        doc.save(str(path))

    def test_missing_output_and_table_mismatch(self):
        with tempfile.TemporaryDirectory() as td:
            out, tpl = Path(td) / "out", Path(td) / "tpl"
            out.mkdir(); tpl.mkdir()
            self._mk_doc(tpl / "a.docx", tables=1)
            self._mk_doc(tpl / "b.docx")
            self._mk_doc(out / "a.docx")  # 缺 b；a 表格数 0 ≠ 模板 1
            ok, issues = m7.check_format(out, tpl)
            self.assertFalse(ok)
            joined = " ".join(issues)
            self.assertIn("缺输出文件", joined)
            self.assertIn("表格数", joined)

    def test_dangling_ref_detected(self):
        with tempfile.TemporaryDirectory() as td:
            out, tpl = Path(td) / "out", Path(td) / "tpl"
            out.mkdir(); tpl.mkdir()
            # 手写一个带悬空 r:id 的 docx zip
            import shutil
            from docx import Document
            doc = Document()
            doc.add_paragraph("正文")
            doc.save(str(out / "a.docx"))
            # 篡改 document.xml：插入 r:id 指向不存在的 rId999
            zpath = out / "a.docx"
            tmp = Path(td) / "tmp.docx"
            with zipfile.ZipFile(str(zpath)) as z:
                xml = z.read("word/document.xml").decode("utf-8")
                xml = xml.replace("</w:body>", '<w:p><w:pPr><w:sectPr><w:footerReference r:id="rId999"/></w:sectPr></w:pPr></w:p></w:body>')
                with zipfile.ZipFile(str(tmp), "w") as zo:
                    for item in z.infolist():
                        data = z.read(item.filename)
                        if item.filename == "word/document.xml":
                            data = xml.encode("utf-8")
                        zo.writestr(item, data)
            tmp.replace(zpath)
            # 模板同文件
            tplf = out / "a.docx"
            shutil.copy2(str(tplf), str(tpl / "a.docx"))
            ok, issues = m7.check_format(out, tpl)
            self.assertFalse(ok)
            self.assertTrue(any("悬空引用" in x and "rId999" in x for x in issues))


class TestCheckSensitive(unittest.TestCase):
    def test_preview_text_flagged(self):
        with tempfile.TemporaryDirectory() as td:
            out, tpl = Path(td) / "out", Path(td) / "tpl"
            out.mkdir(); tpl.mkdir()
            from docx import Document
            doc = Document()
            doc.add_paragraph("此处将插入：营业执照 尺寸：高12cm")
            doc.save(str(out / "a.docx"))
            ok, issues = m7.check_sensitive(out, tpl)
            self.assertFalse(ok)
            self.assertTrue(any("此处将插入" in x for x in issues))


class TestRunChecks(unittest.TestCase):
    def test_report_written(self):
        with tempfile.TemporaryDirectory() as td:
            out, tpl = Path(td) / "out", Path(td) / "tpl"
            out.mkdir(); tpl.mkdir()
            from docx import Document
            for name in ("a.docx", "b.docx"):
                d = Document()
                d.add_paragraph("正文")
                d.save(str(out / name))
                d.save(str(tpl / name))
            res = m7.run_checks(Path(td), Path(td), out_dir=out, tpl_dir=tpl)
            self.assertTrue(res["ok"])
            self.assertTrue((out / "M7检查报告.md").is_file())


if __name__ == "__main__":
    unittest.main()
