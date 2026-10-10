# -*- coding: utf-8 -*-
"""format_applier（V5 格式配置应用器）单元测试。"""
import os
import sys
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "scripts"))
from m5_project import format_applier as fa  # noqa: E402


class ClassifyParaTest(unittest.TestCase):
    def test_chapter(self):
        self.assertEqual(fa.classify_para("四、资格证明及辅助资料表"), "章标题")
        self.assertEqual(fa.classify_para("第四章 投标文件格式"), "章标题")

    def test_first_level(self):
        self.assertEqual(fa.classify_para("表1 组织机构"), "一级标题")
        self.assertEqual(fa.classify_para("表8 拟投入监理人员简历表"), "一级标题")

    def test_sub_title(self):
        self.assertEqual(fa.classify_para("（一）投 标 函"), "小节标题")
        self.assertEqual(fa.classify_para("（1）李林"), "小节标题")
        self.assertEqual(fa.classify_para("（企业业绩）"), "小节标题")
        self.assertEqual(fa.classify_para("企业资质证书"), "小节标题")

    def test_note_and_signoff(self):
        self.assertEqual(fa.classify_para("附：法定代表人有效期内居民身份证正、反面扫描件"), "标注段")
        self.assertEqual(fa.classify_para("法定代表人或其委托代理人（签字或盖章）"), "落款")

    def test_body_and_blank(self):
        self.assertEqual(fa.classify_para("致：和县住房和城乡建设局", ""), "正文")
        self.assertEqual(fa.classify_para("   "), "空行")
        self.assertEqual(fa.classify_para("投标人名称（盖章）：和县建设工程监理有限公司", "开标一览表.docx"), "落款")


class ApplyFormatTest(unittest.TestCase):
    def _mk_doc(self):
        doc = Document()
        p1 = doc.add_paragraph("四、资格证明及辅助资料表")
        p2 = doc.add_paragraph("表1 组织机构")
        p3 = doc.add_paragraph("组 织 机 构")
        p4 = doc.add_paragraph("和县建设工程监理有限公司，经安徽省建设厅批准，成立于2006年11月。")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "企业名称"
        t.cell(0, 1).text = "和县建设工程监理有限公司"
        t.cell(1, 0).text = "法定代表人"
        t.cell(1, 1).text = "邵章华"
        return doc

    def test_apply_format_styles(self):
        doc = self._mk_doc()
        cfg = fa.load_format_config()
        stats = fa.apply_format(doc, "资格证明及辅助资料表.docx", cfg)
        paras = [p for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(len(paras), 4)
        # 章标题 22 居中
        r0 = paras[0].runs[0]
        self.assertEqual(r0.font.size.pt, 22)
        self.assertEqual(paras[0].paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        # 一级标题 16
        self.assertEqual(paras[1].runs[0].font.size.pt, 16)
        # 小节标题 14
        self.assertEqual(paras[2].runs[0].font.size.pt, 14)
        # 正文 12 行距固定28 缩进
        p = paras[3]
        self.assertEqual(p.runs[0].font.size.pt, 12)
        self.assertEqual(p.paragraph_format.line_spacing_rule, WD_LINE_SPACING.EXACTLY)
        self.assertLessEqual(getattr(p.paragraph_format.line_spacing, "pt", 0), 28.0)
        self.assertIsNotNone(p.paragraph_format.first_line_indent)
        self.assertEqual(p.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        # 表格表头加粗
        self.assertTrue(doc.tables[0].cell(0, 0).paragraphs[0].runs[0].font.bold)
        self.assertEqual(stats["表格"], 1)

    def test_cover_no_page_number(self):
        doc = Document()
        doc.add_paragraph("和县2026年老旧小区改造项目（EPC总承包）监理")
        doc.add_paragraph("投 标 文 件")
        doc.add_paragraph("和县建设工程监理有限公司")
        cfg = fa.load_format_config()
        stats = fa.apply_format(doc, "封面.docx", cfg)
        self.assertEqual(stats["页面"], 0)          # 封面无页码


class EnsureOwnPageTest(unittest.TestCase):
    def test_file_level(self):
        doc = Document()
        doc.add_paragraph("开标一览表")
        cfg = fa.load_format_config()
        n = fa.ensure_own_page(doc, "开标一览表.docx", cfg)
        self.assertEqual(n, 1)
        xml = doc.element.body.xml
        self.assertIn("w:type=\"page\"", xml)

    def test_non_own_page_file(self):
        doc = Document()
        doc.add_paragraph("投标函正文")
        cfg = fa.load_format_config()
        n = fa.ensure_own_page(doc, "投标函.docx", cfg)
        self.assertEqual(n, 0)                     # 投标函带括号说明 → 非整文件独占


if __name__ == "__main__":
    unittest.main()
