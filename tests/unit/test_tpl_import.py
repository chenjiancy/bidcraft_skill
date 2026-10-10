# -*- coding: utf-8 -*-
"""unit · M5 企业模板入库（tpl_import：格式契约空白格式提炼）。"""
import os
import sys
import tempfile
import unittest

from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "scripts"))
from m5_project import contract_gen as cg  # noqa: E402
from m5_project import tpl_import as ti    # noqa: E402
from m5_project.gen_blocks import _extract_blocks  # noqa: E402

from tests.unit.test_contract_gen import _build_sample_docx  # noqa: E402


class TestTplImport(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        import json
        import shutil
        cls._tmp = tempfile.mkdtemp(prefix="bidcraft_tplimport_")
        # 样例 docx + 契约（复用契约测试的构造器）
        cls.docx = os.path.join(cls._tmp, "sample.docx")
        _build_sample_docx(cls.docx)
        blocks, _ = _extract_blocks(cls.docx)
        cls.contract = cg.build_contract(blocks, "测试项目")
        # 按 import_templates 路径约定落位：项目级/测试项目/招标解析/格式契约/格式契约.json
        cls.contract_dir = os.path.join(cls._tmp, "项目级", "测试项目", "招标解析", "格式契约")
        os.makedirs(cls.contract_dir)
        cls.contract_path = os.path.join(cls.contract_dir, "格式契约.json")
        with open(cls.contract_path, "w", encoding="utf-8") as f:
            json.dump(cls.contract, f, ensure_ascii=False, indent=2)

    def test_name_map_merged_f06(self):
        """F06a..F06k 合并为单文件（资格证明及辅助资料表）。"""
        ids = ti.TPL_NAME_MAP[7][0]
        self.assertIn("F06a", ids)
        self.assertIn("F06k", ids)

    def test_import_produces_expected_files(self):
        """9 个模板文件生成且内容可打开（跳过 F09/F13/F14）。"""
        tmp_out = os.path.join(self._tmp, "tpl_out")
        res = ti.import_templates(self._tmp, "测试项目", "测试代理",
                                  source=self.docx, out_dir=tmp_out)
        self.assertEqual(len(res["文件"]), 9)
        names = {f["模板"] for f in res["文件"]}
        self.assertIn("模板_封面.docx", names)
        self.assertIn("模板_开标一览表.docx", names)
        self.assertIn("模板_资格证明及辅助资料表.docx", names)
        self.assertIn("模板_承诺函_项目总监到岗.docx", names)
        # F09/F13/F14 未定位 → 跳过（不进入文件列表）
        self.assertTrue(all("中小企业" not in f["模板"] for f in res["文件"]))
        # 每个模板可被 python-docx 打开
        for f in res["文件"]:
            p = os.path.join(tmp_out, f["模板"])
            doc = Document(p)
            self.assertGreater(len(doc.paragraphs) + len(doc.tables), 0)

    def test_import_span_merge(self):
        """资格证明模板范围 = F06a 起点..F06k 终点（附表1-10 连续区间，终点+1=承诺函起点）。"""
        from m5_project.gen_text import _para_text
        import re
        tmp_out = os.path.join(self._tmp, "tpl_out2")
        res = ti.import_templates(self._tmp, "测试项目", "测试代理",
                                  source=self.docx, out_dir=tmp_out)
        blocks, _ = _extract_blocks(self.docx)
        zp = next(f for f in res["文件"] if f["模板"] == "模板_资格证明及辅助资料表.docx")
        cp = next(f for f in res["文件"] if f["模板"] == "模板_承诺函_项目总监到岗.docx")
        lo_txt = re.sub(r"\s+", "", _para_text(blocks[zp["块范围"][0]]["node"]))
        self.assertIn("资格证明及辅助资料表", lo_txt)
        self.assertEqual(zp["块范围"][1] + 1, cp["块范围"][0])


if __name__ == "__main__":
    unittest.main()
