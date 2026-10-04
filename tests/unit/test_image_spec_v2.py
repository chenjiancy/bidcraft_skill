# -*- coding: utf-8 -*-
"""L1 单元测试：M5 图片尺寸统一等比（image_spec v2.0）+ 浮动 anchor 插入。

覆盖：
  - fit_size 统一等比：竖版原图 → 宽=16、高按比例；横版原图 → 宽=16、高收缩；
    身份证保留固定框 8×5（contain）；max_w/max_h 收紧；
  - _to_floating_anchor：inline → anchor（wrapNone、positionH=margin/left、
    positionV=paragraph、元素顺序合法、r:embed 复用）；
  - _insert_image 默认转 anchor、floating=False 保持 inline。
"""
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from m5_project import image_spec as imgsp  # noqa: E402
from m5_project.fill_images import (_insert_image, _to_floating_anchor)  # noqa: E402


def _mkimg(path, w, h, color=(200, 0, 0)):
    from PIL import Image
    im = Image.new("RGB", (w, h), color)
    im.save(str(path))
    return path


class TestFitSize(unittest.TestCase):
    def setUp(self):
        self._td = tempfile.TemporaryDirectory()
        self.vert = _mkimg(Path(self._td.name) / "v.png", 1000, 1400)      # 竖版 1:1.4
        self.horiz = _mkimg(Path(self._td.name) / "h.png", 1400, 950)     # 横版 ~1:0.68
        self.idc = _mkimg(Path(self._td.name) / "idc.png", 960, 600)      # 身份证 1.6:1

    def tearDown(self):
        self._td.cleanup()

    def test_vertical_uniform_ratio(self):
        """竖版原图：宽=16、高按比例（约 22.4）。"""
        w, h = imgsp.fit_size(str(self.vert), "营业执照扫描件")
        self.assertAlmostEqual(w, 16.0, delta=0.1)
        self.assertAlmostEqual(h, 22.4, delta=0.1)

    def test_horizontal_uniform_ratio(self):
        """横版原图：宽=16、高按比例收缩（约 10.9）。"""
        w, h = imgsp.fit_size(str(self.horiz), "职称证")
        self.assertAlmostEqual(w, 16.0, delta=0.1)
        self.assertAlmostEqual(h, 10.9, delta=0.1)

    def test_id_card_fixed_box(self):
        """身份证保留固定框 8×5（contain：960×600 → 8.0×5.0）。"""
        w, h = imgsp.fit_size(str(self.idc), "身份证")
        self.assertAlmostEqual(w, 8.0, delta=0.1)
        self.assertAlmostEqual(h, 5.0, delta=0.1)

    def test_max_w_and_max_h(self):
        """max_w/max_h 收紧：可用宽 16 超 max_h 时按 max_h 定高、宽随比例收缩。"""
        w, h = imgsp.fit_size(str(self.vert), "营业执照扫描件", max_w=15, max_h=None)
        self.assertAlmostEqual(w, 15.0, delta=0.1)
        w2, h2 = imgsp.fit_size(str(self.vert), "营业执照扫描件", max_w=None, max_h=10)
        self.assertAlmostEqual(h2, 10.0, delta=0.1)
        self.assertAlmostEqual(w2, 10.0 / 1.4, delta=0.1)

    def test_unknown_key(self):
        self.assertIsNone(imgsp.fit_size(str(self.vert), "不存在口径"))


class TestFloatingAnchor(unittest.TestCase):
    def test_inline_becomes_anchor(self):
        """_insert_image 默认转 anchor：wrapNone、positionH margin/left、r:embed 复用。"""
        with tempfile.TemporaryDirectory() as td:
            img = _mkimg(Path(td) / "a.png", 1000, 700)
            doc = Document()
            p = doc.add_paragraph()
            _insert_image(p, str(img), 16.0, 11.2)
            drawing = p._p.find(".//" + qn("w:drawing"))
            self.assertIsNotNone(drawing)
            anchor = drawing.find(qn("wp:anchor"))
            self.assertIsNotNone(anchor, "图片应转为 wp:anchor 浮动")
            self.assertIsNotNone(anchor.find(qn("wp:wrapNone")))
            posh = anchor.find(qn("wp:positionH"))
            self.assertEqual(posh.get("relativeFrom"), "margin")
            self.assertEqual(posh.find(qn("wp:align")).text, "left")
            posv = anchor.find(qn("wp:positionV"))
            self.assertEqual(posv.get("relativeFrom"), "paragraph")
            # r:embed 复用：graphic 内 blip 存在
            blip = anchor.find(".//" + qn("a:blip"))
            self.assertIsNotNone(blip)
            self.assertTrue(blip.get(qn("r:embed")))

    def test_floating_false_keeps_inline(self):
        """floating=False（表内图）保持内联。"""
        with tempfile.TemporaryDirectory() as td:
            img = _mkimg(Path(td) / "b.png", 800, 500)
            doc = Document()
            p = doc.add_paragraph()
            _insert_image(p, str(img), 15.0, 9.4, floating=False)
            drawing = p._p.find(".//" + qn("w:drawing"))
            self.assertIsNotNone(drawing)
            self.assertIsNotNone(drawing.find(qn("wp:inline")))
            self.assertIsNone(drawing.find(qn("wp:anchor")))

    def test_to_floating_anchor_element_order(self):
        """anchor 子元素顺序符合 CT_Anchor：simplePos, positionH, positionV,
        extent, effectExtent, wrapNone, docPr, cNvGraphicFramePr, graphic。"""
        from lxml import etree
        with tempfile.TemporaryDirectory() as td:
            img = _mkimg(Path(td) / "c.png", 1000, 800)
            doc = Document()
            p = doc.add_paragraph()
            _insert_image(p, str(img), 16.0, 12.8)
            anchor = p._p.find(".//" + qn("w:drawing")).find(qn("wp:anchor"))
            tags = [etree.QName(ch).localname for ch in anchor]
            expect = ["simplePos", "positionH", "positionV", "extent", "effectExtent",
                      "wrapNone", "docPr", "cNvGraphicFramePr", "graphic"]
            self.assertEqual(tags, expect)


if __name__ == "__main__":
    unittest.main(verbosity=2)
