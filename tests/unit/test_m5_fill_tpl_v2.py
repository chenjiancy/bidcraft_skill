# -*- coding: utf-8 -*-
"""L1 单元测试：fill-tpl v2.0 布局样板范式引擎（占位索引/样板复制/类别聚合/未列不生成）。

覆盖：
  - scan_placeholder_index：通用键/行级键/图片类别分类；简历组内图框不入全局序列
  - aggregate_category：总监识别（职务含总监）+ personnel[0] 兜底；非总监展平；未列返回 []
  - fill_row_templates：人员配备样板行按 personnel 复制、逐行填值；0 人 → 删样板行
  - fill_detail_tables：简历组（标题锚点+简历表+图框段）按 personnel 复制、组内按人员取数不串
  - fill_images_by_category：素材数 < 框数 → 删框+披露；素材数=0（未列）→ 全删+披露
  - fill_generic_text：通用键替换、含行级键段跳过
  - 文本框内占位被扫描与填充
"""
import base64
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))
from docx import Document  # noqa: E402
from docx.enum.text import WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from m5_project import fill_from_tpl as F  # noqa: E402
from m5_project import key_alias          # noqa: E402

# 1×1 PNG（最小合法）
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def _mkimg(path):
    Path(path).write_bytes(_PNG)


def _add_textbox_para(doc, text):
    """在文档末尾段落内嵌文本框，内含 text 段落（合法 w:p>w:r>w:t 结构）。返回文本框内段落对象。"""
    p = doc.add_paragraph()
    run = p.add_run()
    drawing = OxmlElement("w:drawing")
    run._r.append(drawing)
    txbx = OxmlElement("w:txbxContent")
    p_el = OxmlElement("w:p")
    r_el = OxmlElement("w:r")
    t_el = OxmlElement("w:t")
    t_el.text = text
    r_el.append(t_el)
    p_el.append(r_el)
    txbx.append(p_el)
    inline = OxmlElement("wp:inline")
    inline.append(txbx)
    drawing.append(inline)
    from docx.text.paragraph import Paragraph
    return Paragraph(p_el, doc)


def _add_page_break_para(doc):
    """追加分页符段（组边界：简历表后图框段到分页符止）。"""
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


class TestScanIndex(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="bid_fillv2_")
        self.doc_path = os.path.join(self.tmp, "母版.docx")
        doc = Document()
        doc.add_paragraph("项目名称：{{ project_name }}")
        doc.add_paragraph("人员锚点 {{ 人员姓名 }}")
        # 简历表（首行含简历键）
        t = doc.add_table(rows=1, cols=4)
        t.rows[0].cells[0].text = "{{ 人员出生年月 }}"
        t.rows[0].cells[1].text = "{{ 人员学历 }}"
        t.rows[0].cells[2].text = "{{ 人员经历时间 }}"
        t.rows[0].cells[3].text = "{{ 人员经历项目 }}"
        # 简历表后图框段（组内，到分页符止）
        doc.add_paragraph("【图片占位：职称证书】")
        _add_page_break_para(doc)
        # 分页符之后=全局图框
        doc.add_paragraph("【图片占位：社保证明】")
        _add_textbox_para(doc, "【图片占位：岗位证书】")
        doc.save(self.doc_path)

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_classify_keys_and_groups(self):
        doc = Document(self.doc_path)
        idx = F.scan_placeholder_index(doc)
        self.assertIn("project_name", idx["generic_keys"])
        # 简历标题锚点含行级键 → row_paras
        self.assertIn("人员姓名", idx["row_paras"])
        # 简历表（detail_tables resume）
        self.assertEqual([k for _, k in idx["detail_tables"]], ["resume"])
        # 组内图框（职称证书）不入全局
        cats = [c for c, _ in idx["img_ph"]]
        self.assertIn("社保证明", cats)
        self.assertNotIn("职称证书", cats)
        # 文本框内占位被扫描
        self.assertIn("岗位证书", cats)


class TestAggregateCategory(unittest.TestCase):
    def _material(self):
        self.tmp = tempfile.mkdtemp(prefix="bid_agg_")
        p1, p2, p3 = [os.path.join(self.tmp, n) for n in ("a1.png", "a2.png", "a3.png")]
        _mkimg(p1); _mkimg(p2); _mkimg(p3)
        return {
            "personnel": [
                {"name": "张三", "职务": "总监理工程师",
                 "注册证书": [p1], "职称证书": [p2]},
                {"name": "李四", "职务": "专业监理工程师",
                 "职称证书": [p3], "岗位证书": [p1]},
                {"name": "王五", "职务": "监理员", "身份证": [p2]},
            ],
            "social_security_required": [{"paths": [p3]}],
            "performance": [], "qualification_required": [], "honors": [], "iso_certificates": [],
        }

    def tearDown(self):
        import shutil
        if hasattr(self, "tmp"):
            shutil.rmtree(self.tmp, ignore_errors=True)

    def test_director_and_flat(self):
        m = self._material()
        self.assertEqual(F.aggregate_category(m, "总监注册监理工程师证书"),
                         [m["personnel"][0]["注册证书"][0]])
        # 职称证书=非总监展平（李四+王五，王五无 → 只李四）
        self.assertEqual(F.aggregate_category(m, "职称证书"), [m["personnel"][1]["职称证书"][0]])
        # 岗位证书=非总监
        self.assertEqual(F.aggregate_category(m, "岗位证书"), [m["personnel"][1]["岗位证书"][0]])
        # 身份证=全部人员
        self.assertEqual(F.aggregate_category(m, "身份证"), [m["personnel"][2]["身份证"][0]])
        # 社保证明
        self.assertEqual(F.aggregate_category(m, "社保证明"), [m["social_security_required"][0]["paths"][0]])
        # 未列类别 → []
        self.assertEqual(F.aggregate_category(m, "企业法人营业执照"), [])
        self.assertEqual(F.aggregate_category(m, "三体系认证证书"), [])

    def test_director_fallback_first(self):
        m = self._material()
        m["personnel"][0]["职务"] = "专业监理工程师"
        d = F._find_director(m["personnel"])
        self.assertEqual(d["name"], "张三")  # 职务无总监 → personnel[0] 兜底


class TestRowTemplateReplicate(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="bid_rows_")
        self.p1 = os.path.join(self.tmp, "img1.png")
        _mkimg(self.p1)

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _doc_with_row(self, path):
        doc = Document()
        doc.add_paragraph("开标")
        t = doc.add_table(rows=2, cols=3)
        t.rows[0].cells[0].text = "姓名"
        t.rows[0].cells[1].text = "年龄"
        t.rows[0].cells[2].text = "职务"
        t.rows[1].cells[0].text = "{{ 人员姓名 }}"
        t.rows[1].cells[1].text = "{{ 人员年龄 }}"
        t.rows[1].cells[2].text = "{{ 人员职务 }}"
        doc.save(path)
        return doc

    def test_replicate_and_fill(self):
        path = os.path.join(self.tmp, "m.docx")
        doc = self._doc_with_row(path)
        mat = {"personnel": [
            {"name": "张三", "年龄": "30", "职务": "总监理工程师"},
            {"name": "李四", "年龄": "28", "职务": "专业监理工程师"},
            {"name": "王五", "年龄": "26", "职务": "监理员"},
        ]}
        idx = F.scan_placeholder_index(doc)
        F.fill_row_templates(doc, idx, mat, Path(self.tmp))
        t = doc.tables[0]
        self.assertEqual(len(t.rows), 4)  # 表头 + 3 人
        names = [t.rows[i].cells[0].text for i in (1, 2, 3)]
        self.assertEqual(names, ["张三", "李四", "王五"])
        self.assertEqual(t.rows[1].cells[2].text, "总监理工程师")

    def test_zero_staff_removes_row(self):
        path = os.path.join(self.tmp, "m.docx")
        doc = self._doc_with_row(path)
        mat = {"personnel": []}
        idx = F.scan_placeholder_index(doc)
        F.fill_row_templates(doc, idx, mat, Path(self.tmp))
        t = doc.tables[0]
        self.assertEqual(len(t.rows), 1)  # 只剩表头，样板行删除（未列不生成）


class TestResumeGroupReplicate(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="bid_resume_")
        self.z1 = os.path.join(self.tmp, "z1.png")
        self.z2 = os.path.join(self.tmp, "z2.png")
        _mkimg(self.z1); _mkimg(self.z2)

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _doc(self, path):
        doc = Document()
        doc.add_paragraph("人员锚点 {{ 人员姓名 }}")
        t = doc.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "{{ 人员出生年月 }}"
        t.rows[0].cells[1].text = "{{ 人员经历时间 }}"
        doc.add_paragraph("【图片占位：职称证书】")
        _add_page_break_para(doc)  # 组边界
        doc.add_paragraph("【图片占位：社保证明】")  # 分页符后=全局
        doc.save(path)
        return doc

    def test_two_people_two_groups_no_cross(self):
        path = os.path.join(self.tmp, "m.docx")
        doc = self._doc(path)
        mat = {"personnel": [
            {"name": "张三", "出生年月": "1980.1", "经历": [{"时间": "2010-2013", "项目": "A"}],
             "职称证书": [self.z1]},
            {"name": "李四", "出生年月": "1985.2", "经历": [],
             "职称证书": [self.z2]},
        ]}
        idx = F.scan_placeholder_index(doc)
        missing = []
        F.fill_detail_tables(doc, idx, mat, Path(self.tmp), missing)
        self.assertEqual(missing, [])
        # 两组：锚点标题 ×2、简历表 ×2、图框段 ×2
        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(texts.count("人员锚点 张三"), 1)
        self.assertEqual(texts.count("人员锚点 李四"), 1)
        # 表格 ×2，各自填对应人员
        self.assertEqual(len(doc.tables), 2)
        t0 = "".join(c.text for c in doc.tables[0].rows[0].cells)
        t1 = "".join(c.text for c in doc.tables[1].rows[0].cells)
        self.assertIn("1980.1", t0)
        self.assertIn("1985.2", t1)
        # 图框段插入图片（drawing 数=2：每人职称证书各 1）
        n = len(doc.element.body.findall(".//" + qn("w:drawing")))
        self.assertEqual(n, 2)

    def test_zero_people_removes_group(self):
        path = os.path.join(self.tmp, "m.docx")
        doc = self._doc(path)
        mat = {"personnel": []}
        idx = F.scan_placeholder_index(doc)
        missing = []
        F.fill_detail_tables(doc, idx, mat, Path(self.tmp), missing)
        self.assertEqual(len(doc.tables), 0)  # 整组删除（未列不生成）


class TestImagesByCategory(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="bid_imgcat_")
        self.s1 = os.path.join(self.tmp, "s1.png")
        self.s2 = os.path.join(self.tmp, "s2.png")
        _mkimg(self.s1); _mkimg(self.s2)

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_one_material_two_frames(self):
        doc = Document()
        doc.add_paragraph("【图片占位：社保证明】")
        doc.add_paragraph("【图片占位：社保证明】")
        path = os.path.join(self.tmp, "m.docx")
        doc.save(path)
        doc = Document(path)
        mat = {"social_security_required": [{"paths": [self.s1]}]}
        idx = F.scan_placeholder_index(doc)
        missing = []
        F.fill_images_by_category(doc, idx, mat, Path(self.tmp), missing)
        # 第 1 框插入，第 2 框删框+披露
        n = len(doc.element.body.findall(".//" + qn("w:drawing")))
        self.assertEqual(n, 1)
        self.assertTrue(any("素材不足" in note or "未列" in note for _, note in missing))
        # 占位文案全部清除（插图段删占位文案，不足框已删段）
        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(sum(1 for t in texts if "【图片占位：" in t), 0)

    def test_not_listed_removes_all(self):
        doc = Document()
        doc.add_paragraph("【图片占位：三体系认证证书】")
        path = os.path.join(self.tmp, "m.docx")
        doc.save(path)
        doc = Document(path)
        mat = {"iso_certificates": []}
        idx = F.scan_placeholder_index(doc)
        missing = []
        F.fill_images_by_category(doc, idx, mat, Path(self.tmp), missing)
        self.assertEqual(len(doc.element.body.findall(".//" + qn("w:drawing"))), 0)
        self.assertTrue(any("素材清单未列" in note for _, note in missing))


class TestAiImgCategory(unittest.TestCase):
    """商务标红线（v2.5）：fill-tpl（商务标引擎）**禁用 AI 配图**——
    『【图片占位：AI配图-<section>】』一律删框+披露（素材清单无则不生成），
    即使配图清单有 saved 项也不消费；m_img 管道仅服务非商务标场景。
    """

    def _mk_plan(self, items):
        import json
        d = Path(self.tmp) / "配图"
        d.mkdir(parents=True, exist_ok=True)
        (d / "配图清单_冒烟.json").write_text(
            json.dumps({"project": "冒烟", "items": items}, ensure_ascii=False), encoding="utf-8")

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="bid_aiimg_")

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_saved_not_embedded_removed(self):
        # 商务标红线：清单即使有 saved 项也不嵌入——AI配图 占位全部删框+披露
        self._mk_plan([
            {"id": "img_001", "section": "项目概况与监理工作内容", "status": "saved",
             "file": "配图/a.png", "ratio": "16:9"},
            {"id": "img_002", "section": "评分办法", "status": "pending", "file": ""},
        ])
        _mkimg(os.path.join(self.tmp, "配图", "a.png"))
        doc = Document()
        doc.add_paragraph("【图片占位：AI配图-项目概况与监理工作内容】")
        doc.add_paragraph("【图片占位：AI配图-评分办法】")
        path = os.path.join(self.tmp, "m.docx")
        doc.save(path)
        doc = Document(path)
        idx = F.scan_placeholder_index(doc)
        missing = []
        F.fill_images_by_category(doc, idx, {}, Path(self.tmp), missing)
        n = len(doc.element.body.findall(".//" + qn("w:drawing")))
        self.assertEqual(n, 0)  # saved 也不嵌入：商务标禁 AI 配图
        self.assertEqual(len(missing), 2)  # 两个占位均披露
        self.assertTrue(all("商务标禁 AI 配图" in note for _, note in missing))
        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(sum(1 for t in texts if "【图片占位：" in t), 0)

    def test_no_plan_removes_all(self):
        doc = Document()
        doc.add_paragraph("【图片占位：AI配图-项目概况】")
        path = os.path.join(self.tmp, "m.docx")
        doc.save(path)
        doc = Document(path)
        idx = F.scan_placeholder_index(doc)
        missing = []
        F.fill_images_by_category(doc, idx, {}, Path(self.tmp), missing)
        self.assertEqual(len(doc.element.body.findall(".//" + qn("w:drawing"))), 0)
        self.assertTrue(any("AI配图-项目概况" in c for c, _ in missing))

    def test_spec_key_prefix(self):
        self.assertEqual(F._spec_key("AI配图-项目概况与监理工作内容"), "AI配图")
        self.assertEqual(F._spec_key("AI配图-评分办法"), "AI配图")
        # 既有映射不受影响
        self.assertEqual(F._spec_key("企业法人营业执照"), "营业执照扫描件")
        self.assertEqual(F._spec_key("总监注册监理工程师证书"), "注册监理工程师证书")


class TestGenericText(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="bid_gen_")

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_generic_replaced_row_skipped(self):
        doc = Document()
        doc.add_paragraph("项目名称：{{ project_name }}")
        doc.add_paragraph("锚点 {{ 人员姓名 }}")
        path = os.path.join(self.tmp, "m.docx")
        doc.save(path)
        doc = Document(path)
        n = F.fill_generic_text(doc, {"project_name": "测试项目"})
        self.assertEqual(n, 1)
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("项目名称：测试项目", texts)
        # 行级键段保持不动
        self.assertTrue(any("{{ 人员姓名 }}" in t for t in texts))

    def test_unknown_keys_reported(self):
        """v2.4 白名单：未知键披露（不静默）。"""
        doc = Document()
        doc.add_paragraph("项目：{{ project_name }}")
        doc.add_paragraph("拼错键：{{ project_name_typo }}")
        doc.add_paragraph("行级键合法：{{ 人员姓名 }}")
        path = os.path.join(self.tmp, "u.docx")
        doc.save(path)
        doc = Document(path)
        self.assertEqual(F.scan_unknown_keys(doc, key_alias.registered_keys()),
                         ["project_name_typo"])

    def test_alias_key_resolution(self):
        """v2.4 别名解析：模板写中文别名（{{ 企业名称 }}）也能填充。"""
        from m5_project import key_alias
        doc = Document()
        doc.add_paragraph("投标人：{{ 企业名称 }}")
        path = os.path.join(self.tmp, "a.docx")
        doc.save(path)
        doc = Document(path)
        v = {"company_name": "和县建设工程监理有限公司"}
        for mk, spec in key_alias.KEY_SPEC.items():
            al = spec.get("别名")
            if al and mk in v:
                v.setdefault(al, v[mk])
        n = F.fill_generic_text(doc, v)
        self.assertEqual(n, 1)
        self.assertIn("投标人：和县建设工程监理有限公司", [p.text for p in doc.paragraphs])


if __name__ == "__main__":
    unittest.main(verbosity=2)
