# -*- coding: utf-8 -*-
"""L2 单元测试：M6 模板部件库拆解器（partlib.scan_master）+ 部件库定位。

覆盖：docx 母版结构识别（cover/toc/body/closing）、内容参数化（企业名/项目名/日期→占位）、
部件 JSON schema、部件库目录兜底（仓库共享资源）。"""
import json
import tempfile
import unittest
from pathlib import Path

import sys
import os

_REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(_REPO / "scripts"))

from m6_tech import partlib  # noqa: E402

ENT = "安徽安硕工程咨询有限公司"
PROJ = "滨江智慧园区机电安装工程"


def _mk_master(path):
    """迷你 docx 母版：封面（项目名/标题/企业名/日期）→ 目录 → 正文 → 结尾。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    doc = Document()
    for text, size in [(PROJ, 26), ("投标文件（技术标）", 16), (ENT, 14), ("2026年10月24日", 12)]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size = Pt(size)
    doc.add_paragraph("目　录")
    doc.add_paragraph("第一章　工程概况")
    doc.add_paragraph("按总进度计划组织施工……")
    doc.add_paragraph("（以下无正文）")
    doc.save(str(path))
    return str(path)


class TestScanMaster(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.docx = _mk_master(str(Path(self.tmp.name) / "母版.docx"))

    def tearDown(self):
        self.tmp.cleanup()

    def test_structure_recognition_and_param(self):
        doc = partlib.scan_master(self.docx, ENT, PROJ)
        self.assertEqual(doc["schema"], partlib.PART_SCHEMA)
        self.assertEqual(doc["source"], "母版.docx")
        parts = doc["parts"]
        # cover：toc 前 4 段，企业名/项目名/日期均已参数化
        cover = parts["cover"]
        self.assertIsNotNone(cover)
        texts = [p["text"] for p in cover["paragraphs"]]
        self.assertEqual(texts[0], "{{ project_name }}")
        self.assertEqual(texts[1], "投标文件（技术标）")
        self.assertEqual(texts[2], "{{ enterprise_name }}")
        self.assertEqual(texts[3], "{{ bid_date }}")
        self.assertTrue(cover.get("pageBreak"))
        # 样式：26pt 居中
        self.assertEqual(cover["paragraphs"][0]["style"]["size"], 26)
        self.assertEqual(cover["paragraphs"][0]["style"]["align"], "center")
        # toc 识别
        self.assertIsNotNone(parts["toc"])
        self.assertEqual(parts["toc"]["title"], "目　录")
        # colophon 不自动切分
        self.assertIsNone(parts["colophon"])
        # body 样式基线（正文段 12pt）
        self.assertEqual(parts["body"]["font"]["size"], 12)
        # closing 识别
        closing = parts["closing"]
        self.assertIsNotNone(closing)
        self.assertTrue(any("以下无正文" in p["text"] for p in closing["paragraphs"]))

    def test_scan_without_enterprise_keeps_text(self):
        doc = partlib.scan_master(self.docx, "", "")
        cover = doc["parts"]["cover"]
        self.assertEqual(cover["paragraphs"][2]["text"], ENT)  # 未参数化（企业名保留原文）
        # 日期为通用占位：不依赖企业/项目，仍参数化
        self.assertEqual(cover["paragraphs"][3]["text"], "{{ bid_date }}")


class TestPartLibDir(unittest.TestCase):
    def test_repo_fallback(self):
        import argparse
        # 不存在的根：仍能落到仓库共享资源部件库（specimen 内置样例等场景）
        a = argparse.Namespace(root=str(Path(tempfile.mkdtemp(prefix="bid_none_"))))
        d = partlib.part_lib_dir(a)
        self.assertIsNotNone(d)
        self.assertTrue(d.is_dir())

    def test_part_schema_roundtrip(self):
        doc = {
            "schema": partlib.PART_SCHEMA,
            "name": "x",
            "parts": {"cover": {"paragraphs": [], "pageBreak": True}},
        }
        s = json.dumps(doc, ensure_ascii=False)
        back = json.loads(s)
        self.assertEqual(back["schema"], "m6.part/v1")


if __name__ == "__main__":
    unittest.main()
