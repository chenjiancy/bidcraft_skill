# -*- coding: utf-8 -*-
"""L1 单元测试：M2 企业级模板库占位预览图集（template_preview）。

覆盖 _resolve_assets（精确/关键词兜底/组合待补）、_preview_box_size
（精确/兜底/默认）与 build_template_previews 端到端（临时企业目录：
预览副本生成、预览框 docPr 打标、原文件 hash 不变、说明文件产出）。
"""
import hashlib
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))
from m2_template import template_preview as TP  # noqa: E402


def _have_docx():
    try:
        import docx  # noqa: F401
        return True
    except Exception:
        return False


class TestResolveAssets(unittest.TestCase):
    def setUp(self):
        self._td = tempfile.TemporaryDirectory()
        self.lib = Path(self._td.name) / "素材库"
        d = self.lib / "资质"
        d.mkdir(parents=True)
        (d / "营业执照_副本_长期.png").write_bytes(b"\x89PNG-fake")
        (self.lib / "财务" / "财务证照").mkdir(parents=True)
        (self.lib / "财务" / "财务证照" / "开户许可证_长期.png").write_bytes(b"\x89PNG-fake")

    def tearDown(self):
        self._td.cleanup()

    def test_exact_single_hit(self):
        """FILL_IMG_MAP 精确单图且素材存在 → 1 个素材。"""
        a = TP._resolve_assets("【图片：企业法人营业执照（副本）扫描件】", self.lib)
        self.assertIsNotNone(a)
        self.assertEqual(len(a), 1)
        self.assertTrue(Path(a[0][0]).is_file())

    def test_exact_single_missing_file(self):
        """精确键但素材库缺文件 → None（待补）。"""
        a = TP._resolve_assets("【图片：企业基本账户开户许可证扫描件】", self.lib / "空目录")
        self.assertIsNone(a)

    def test_keyword_fallback_hit(self):
        """模板库变体文案（基本存款账户信息）经关键词兜底命中开户许可证。"""
        a = TP._resolve_assets("【图片：基本账户开户许可证（或基本存款账户信息）扫描件】", self.lib)
        self.assertIsNotNone(a)
        self.assertTrue("开户许可证" in Path(a[0][0]).name)

    def test_combo_returns_none(self):
        """组合/项目相关占位（无固定素材）→ None（按项目素材清单填充）。"""
        self.assertIsNone(TP._resolve_assets("【图片：三体系认证证书】", self.lib))
        self.assertIsNone(TP._resolve_assets("【图片：业绩项目一证明材料】", self.lib))


class TestPreviewBoxSize(unittest.TestCase):
    def test_exact_spec(self):
        self.assertEqual(TP._preview_box_size("【图片：三体系认证证书】"), (16.0, 23.0))
        self.assertEqual(TP._preview_box_size("【图片：法定代表人身份证正、反面扫描件】"), (8.0, 5.0))

    def test_keyword_fallback(self):
        self.assertEqual(TP._preview_box_size("【图片：人员职称证书】"), (16.0, 23.0))
        self.assertEqual(TP._preview_box_size("【图片：监理示范工程证书】"), (14.0, 11.0))

    def test_default_fallback(self):
        self.assertEqual(TP._preview_box_size("【图片：业绩项目一证明材料】"), (16.0, 12.0))


@unittest.skipUnless(_have_docx(), "依赖 python-docx")
class TestBuildPreviews(unittest.TestCase):
    def test_end_to_end(self):
        """临时企业目录：预览副本生成 + 预览框打标 + 原文件不变 + 说明产出。"""
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            ent = Path(td)
            tdir = ent / "企业级" / "模板库" / "某代理" / "投标"
            lib = ent / "企业级" / "素材库"
            (tdir).mkdir(parents=True)
            (lib / "资质").mkdir(parents=True)
            (lib / "资质" / "营业执照_副本_长期.png").write_bytes(b"\x89PNG-fake")
            src = tdir / "资格证明.docx"
            doc = Document()
            doc.add_paragraph("【图片：企业法人营业执照（副本）扫描件】")
            doc.add_paragraph("【图片：三体系认证证书】")
            doc.save(str(src))
            h_before = hashlib.sha256(src.read_bytes()).hexdigest()

            res = TP.build_template_previews(ent, "某代理", "投标")
            out = Path(res["输出目录"])
            prev = out / "预览_资格证明.docx"
            self.assertTrue(prev.is_file())
            self.assertEqual(res["合计"]["图片占位"], 2)
            self.assertEqual(res["合计"]["有素材"], 1)
            self.assertEqual(res["合计"]["待补"], 1)
            # 预览框 docPr 标记
            d2 = Document(str(prev))
            marks = []
            for p in d2.paragraphs:
                for r in p.runs:
                    for el in r._r.iter():
                        if el.tag.endswith("}docPr") and el.get("descr"):
                            marks.append(el.get("descr"))
            self.assertEqual(len(marks), 2)
            # 原文件未被修改
            self.assertEqual(hashlib.sha256(src.read_bytes()).hexdigest(), h_before)
            # 说明文件
            self.assertTrue(Path(res["说明文件"]).is_file())


if __name__ == "__main__":
    unittest.main(verbosity=2)
