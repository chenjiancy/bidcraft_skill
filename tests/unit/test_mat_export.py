# -*- coding: utf-8 -*-
"""L1 · mat_export（素材清单组装导出：v5.1 表 → 生成器结构）。"""
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))

from m5_project import mat_export as me

ENT = "test-ent"
PROJ = "示例项目（EPC总承包）监理"


def _mk_project(root, rows=None):
    """构造临时项目目录 + v5.1 表 + 素材库（台账+文件）。返回 (ent, proj_dir)。"""
    ent = root / ENT
    proj_dir = ent / "项目级" / PROJ
    tdir = proj_dir / "招标解析"
    tdir.mkdir(parents=True)
    (ent / "企业级" / "素材库" / "资质").mkdir(parents=True)
    (ent / "企业级" / "素材库" / "荣誉").mkdir(parents=True)
    (ent / "企业级" / "素材库" / "业绩" / "示例业绩A_20230101_张三").mkdir(parents=True)
    (ent / "企业级" / "素材库" / "人员" / "李四" / "个人荣誉").mkdir(parents=True)
    (ent / "企业级" / "素材库" / "人员" / "王五" / "个人荣誉").mkdir(parents=True)
    (ent / "企业级" / "素材库" / "人员" / "李四" / "注册证书").mkdir(parents=True)
    (proj_dir / "项目资料").mkdir(parents=True)
    # 素材文件
    (ent / "企业级" / "素材库" / "资质" / "营业执照_副本_长期.png").write_bytes(b"x")
    (ent / "企业级" / "素材库" / "资质" / "房屋建筑工程监理甲级_20281222_P1.png").write_bytes(b"x")
    (ent / "企业级" / "素材库" / "资质" / "ISO9001_20290318.png").write_bytes(b"x")
    (ent / "企业级" / "素材库" / "荣誉" / "监理示范工程_20241201.png").write_bytes(b"x")
    (ent / "企业级" / "素材库" / "人员" / "李四" / "个人荣誉" / "优秀总监理工程师_20231201.png").write_bytes(b"x")
    (ent / "企业级" / "素材库" / "人员" / "王五" / "个人荣誉" / "优秀总监理工程师_20231201.png").write_bytes(b"x")
    (ent / "企业级" / "素材库" / "人员" / "李四" / "注册证书" / "监理工程师_20280123_20270220.png").write_bytes(b"x")
    (ent / "企业级" / "素材库" / "业绩" / "示例业绩A_20230101_张三" / "监理合同_P0.png").write_bytes(b"x")
    (proj_dir / "项目资料" / "社保单位参保证明_示例公司_20260807.png").write_bytes(b"x")
    # 台账
    ledger = [
        {"id": "a", "category": "资质", "rel_path": "资质/营业执照_副本_长期.png"},
        {"id": "b", "category": "资质", "rel_path": "资质/房屋建筑工程监理甲级_20281222_P1.png"},
        {"id": "c", "category": "资质", "rel_path": "资质/ISO9001_20290318.png"},
        {"id": "d", "category": "荣誉", "rel_path": "荣誉/监理示范工程_20241201.png"},
        {"id": "e", "category": "人员", "rel_path": "人员/李四/个人荣誉/优秀总监理工程师_20231201.png"},
        {"id": "f", "category": "人员", "rel_path": "人员/王五/个人荣誉/优秀总监理工程师_20231201.png"},
        {"id": "g", "category": "人员", "rel_path": "人员/李四/注册证书/监理工程师_20280123_20270220.png"},
        {"id": "h", "category": "业绩", "rel_path": "业绩/示例业绩A_20230101_张三/监理合同_P0.png"},
    ]
    (ent / "企业级" / "素材库" / "素材台账.json").write_text(
        json.dumps(ledger, ensure_ascii=False), encoding="utf-8")
    # v5.1 表（素材清单 sheet：R01/R03/R10/R13/R14/R19 行 + 附加素材行）
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "素材清单"
    rows = rows or [
        # (行, A需求编号, G素材, K确认, H查询结果)
        (3, "R01", "营业执照_副本_长期.png", "✓ 已确认", "待查询"),
        (4, "R02", "房屋建筑工程监理甲级_20281222_P1.png", "✓ 已确认", "待查询"),
        (6, "R03", "ISO9001_20290318.png", "✓ 已确认", "待查询"),
        (9, "R04", "李四：监理工程师_20280123_20270220.png", "✓ 已确认", "待查询"),
        (16, "R09", "李四：简历_20261001_P0.png", "✓ 已确认", "待查询"),
        (21, "R10", "示例业绩A_20230101_张三\\监理合同（1页）", "✓ 已确认", "待查询"),
        (26, "R13", "监理示范工程_20241201.png", "✓ 已确认", "待查询"),
        (28, "R14", "李四：优秀总监理工程师_20231201.png", "✓ 已确认", "待查询"),
        (37, "R19", "姓名", "✓ 已确认", "待查询"),
    ]
    for r, a, g, k, _h in rows:
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=7, value=g)
        ws.cell(row=r, column=11, value=k)
    ws.cell(row=37, column=8, value="测试代理人")
    ws.cell(row=38, column=7, value="手机号")
    ws.cell(row=38, column=8, value="13900000000")
    wb.save(tdir / "素材清单_空白表_v5.1.xlsm")
    return ent, proj_dir


class ResolveAssetTest(unittest.TestCase):
    def setUp(self):
        self._tmp = Path(tempfile.mkdtemp(prefix="bidcraft_mat_"))
        self.ent, self.proj_dir = _mk_project(self._tmp)
        self.lib = self.ent / "企业级" / "素材库"
        self.ledger = json.loads((self.lib / "素材台账.json").read_text(encoding="utf-8"))

    def tearDown(self):
        shutil.rmtree(str(self._tmp), ignore_errors=True)

    def test_纯文件名台账匹配(self):
        got = me._resolve_asset("营业执照_副本_长期.png", self.lib, self.ledger).replace("\\", "/")
        self.assertTrue(got.endswith("资质/营业执照_副本_长期.png"))

    def test_姓名前缀过滤同名文件(self):
        """「李四：优秀总监理工程师」须命中 人员/李四/ 而非 王五 同名文件。"""
        got = me._resolve_asset("李四：优秀总监理工程师_20231201.png", self.lib, self.ledger).replace("\\", "/")
        self.assertIn("人员/李四/", got)
        self.assertNotIn("王五", got)

    def test_业绩目录段定位(self):
        got = me._resolve_asset("示例业绩A_20230101_张三\\监理合同（1页）", self.lib, self.ledger).replace("\\", "/")
        self.assertTrue(got.endswith("业绩/示例业绩A_20230101_张三/监理合同_P0.png"))

    def test_无姓名前缀取首个同名(self):
        got = me._resolve_asset("优秀总监理工程师_20231201.png", self.lib, self.ledger)
        self.assertIn("个人荣誉", got)

    def test_解析失败返回空串(self):
        self.assertEqual(me._resolve_asset("不存在的素材_999999.png", self.lib, self.ledger), "")

    def test_空值返回空串(self):
        self.assertEqual(me._resolve_asset("", self.lib, self.ledger), "")


class ExportMaterialTest(unittest.TestCase):
    def setUp(self):
        self._tmp = Path(tempfile.mkdtemp(prefix="bidcraft_mat_"))
        self.ent, self.proj_dir = _mk_project(self._tmp)

    def tearDown(self):
        shutil.rmtree(str(self._tmp), ignore_errors=True)

    def test_sections组装(self):
        m = me.export_material(self.ent, PROJ, social_path="")
        self.assertEqual(len(m["qualification_required"]), 2)      # R01 营业执照 + R02 房建甲级
        self.assertEqual(len(m["iso_certificates"]), 1)
        self.assertEqual(len(m["honors"]), 2)                      # R13 示范工程 + R14 总监荣誉
        self.assertEqual(len(m["performance"]), 1)                 # R10 企业业绩
        for it in m["qualification_required"] + m["iso_certificates"] \
                + m["honors"] + m["performance"]:
            self.assertTrue(it["path"], "path 应为非空：%s" % it)

    def test_personnel挂载与角色(self):
        m = me.export_material(self.ent, PROJ, social_path="")
        persons = {p["name"]: p for p in m["personnel"]}
        self.assertIn("李四", persons)
        self.assertEqual(len(persons["李四"]["注册证书"]), 1)
        self.assertTrue(persons["李四"]["注册证书"][0].endswith("监理工程师_20280123_20270220.png"))

    def test_企业基础信息与配置口径(self):
        m = me.export_material(self.ent, PROJ, social_path="")
        self.assertEqual(m["企业基础信息"]["法定代表人姓名"], "邵章华")
        self.assertEqual(m["企业基础信息"]["委托代理人姓名"], "测试代理人")
        self.assertEqual(m["企业基础信息"]["委托代理人手机"], "13900000000")
        self.assertEqual(m["监理人员配置口径"]["项目总监"], 1)
        self.assertIn("投标资格专业", m)

    def test_items保留M4需求清单(self):
        # 预置 M4 产物（items 需求清单）
        (self.proj_dir / "招标解析" / ("素材清单_%s.json" % PROJ)).write_text(
            json.dumps({"project": PROJ, "items": [{"category": "资质", "purpose": "x"}]},
                       ensure_ascii=False), encoding="utf-8")
        m = me.export_material(self.ent, PROJ, social_path="")
        self.assertEqual(len(m["items"]), 1)

    def test_social_path传入(self):
        m = me.export_material(self.ent, PROJ, social_path="E:/社保.png")
        self.assertEqual(m["social_security_required"][0]["path"], "E:/社保.png")


class WriteMaterialTest(unittest.TestCase):
    def setUp(self):
        self._tmp = Path(tempfile.mkdtemp(prefix="bidcraft_mat_"))
        self.ent, self.proj_dir = _mk_project(self._tmp)

    def tearDown(self):
        shutil.rmtree(str(self._tmp), ignore_errors=True)

    def test_落盘默认路径(self):
        m = me.export_material(self.ent, PROJ, social_path="")
        p = me.write_material(self.ent, PROJ, m)
        self.assertTrue(p.is_file())
        d = json.loads(p.read_text(encoding="utf-8"))
        self.assertEqual(d["project"], PROJ)
        self.assertIn("personnel", d)
        self.assertIn("qualification_required", d)


if __name__ == "__main__":
    unittest.main(verbosity=2)
