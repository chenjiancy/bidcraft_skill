# -*- coding: utf-8 -*-
"""M7 T17 技术标检查扩充单测：图表位置 / 空白页 / 引用一致性 + M10 补救联动。

docx 用 python-docx 合成（含 inline 图片段 w:drawing、分页符 w:br type=page、空段）。
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

from docx.oxml import OxmlElement       # noqa: E402
from docx.oxml.ns import qn             # noqa: E402

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "scripts"))

from m7_check import rule_check as R        # noqa: E402
from m10_remedy import remedy               # noqa: E402


def _add_img(doc, caption):
    p = doc.add_paragraph(caption)
    run = p.add_run()
    drawing = OxmlElement("w:drawing")
    inline = OxmlElement("wp:inline")
    inline.append(OxmlElement("wp:extent"))
    drawing.append(inline)
    run._r.append(drawing)


def _mk_docx(path, body=20, image_mid=False, image_appendix=False,
             page_break_empty=False, empty_run=0, names=None):
    """合成 docx：正文段 + 可选图段/分页空段/连续空段/指定名称文本。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph("封面：技术标")
    if image_mid:
        # 图段插在正文中段（前后都有正文）
        for i in range(body // 2):
            doc.add_paragraph("监理正文段落%d：全过程监理。" % i)
        _add_img(doc, "图 1：工程平面示意")
        for i in range(body // 2, body):
            doc.add_paragraph("监理正文段落%d：全过程监理。" % i)
    else:
        for i in range(body):
            doc.add_paragraph("监理正文段落%d：全过程监理。" % i)
        if image_appendix:
            _add_img(doc, "附图：平面图")
    if page_break_empty:
        doc.add_paragraph("")                       # 前置空段
        p = doc.add_paragraph()
        run = p.add_run()
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run._r.append(br)
        doc.add_paragraph("")                       # 后置空段（双空夹分页）
    for _ in range(empty_run):
        doc.add_paragraph("")
    if names:
        doc.add_paragraph(names)
    doc.add_paragraph("结尾段：本技术标制作完成。")
    doc.save(str(path))


class TechFigureTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.proj = Path(self.tmp.name) / "P"

    def tearDown(self):
        self.tmp.cleanup()

    def test_no_image_passes(self):
        f = self.proj / "t.docx"
        self.proj.mkdir(parents=True)
        _mk_docx(f)
        ok, issues = R.check_tech_figure_placement(str(f), str(self.proj), {}, {})
        self.assertTrue(ok, issues)

    def test_mid_text_image_flagged(self):
        f = self.proj / "t.docx"
        self.proj.mkdir(parents=True)
        _mk_docx(f, body=10, image_mid=True)
        ok, issues = R.check_tech_figure_placement(str(f), str(self.proj), {}, {})
        self.assertFalse(ok)
        self.assertTrue(any("插入图表" in x for x in issues))

    def test_appendix_image_passes(self):
        f = self.proj / "t.docx"
        self.proj.mkdir(parents=True)
        _mk_docx(f, body=10, image_appendix=True)
        ok, issues = R.check_tech_figure_placement(str(f), str(self.proj), {}, {})
        self.assertTrue(ok, issues)


class TechBlankPagesTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.proj = Path(self.tmp.name) / "P"

    def tearDown(self):
        self.tmp.cleanup()

    def test_clean_passes(self):
        f = self.proj / "t.docx"
        self.proj.mkdir(parents=True)
        _mk_docx(f)
        ok, issues = R.check_tech_blank_pages(str(f), str(self.proj), {}, {})
        self.assertTrue(ok, issues)

    def test_page_break_empty_flagged(self):
        f = self.proj / "t.docx"
        self.proj.mkdir(parents=True)
        _mk_docx(f, page_break_empty=True)
        ok, issues = R.check_tech_blank_pages(str(f), str(self.proj), {}, {})
        self.assertFalse(ok)
        self.assertTrue(any("分页符" in x for x in issues))

    def test_single_page_break_after_text_passes(self):
        """真实导出结构：章末/TOC 后分页符空段两侧有正文 → 不误报为空白页。"""
        f = self.proj / "t.docx"
        self.proj.mkdir(parents=True)
        from docx import Document
        doc = Document()
        doc.add_paragraph("封面：技术标")
        doc.add_paragraph("目录：第1章 工程概况")
        p = doc.add_paragraph()
        run = p.add_run()
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run._r.append(br)
        doc.add_paragraph("第一章 工程概况与监理范围")
        doc.save(str(f))
        ok, issues = R.check_tech_blank_pages(str(f), str(self.proj), {}, {})
        self.assertTrue(ok, issues)

    def test_three_empty_paras_flagged(self):
        f = self.proj / "t.docx"
        self.proj.mkdir(parents=True)
        _mk_docx(f, empty_run=3)
        ok, issues = R.check_tech_blank_pages(str(f), str(self.proj), {}, {})
        self.assertFalse(ok)
        self.assertTrue(any("空段" in x for x in issues))

    def test_two_empty_paras_passes(self):
        f = self.proj / "t.docx"
        self.proj.mkdir(parents=True)
        _mk_docx(f, empty_run=2)
        ok, issues = R.check_tech_blank_pages(str(f), str(self.proj), {}, {})
        self.assertTrue(ok, issues)


class TechRefConsistencyTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.proj = self.root / "P"
        self.proj.mkdir(parents=True)
        (self.proj / "技术标").mkdir(parents=True)
        self.f = self.proj / "t.docx"

    def tearDown(self):
        self.tmp.cleanup()

    def _settings(self, bid_mode):
        (self.proj / "技术标" / "设置.json").write_text(
            json.dumps({"schema": "m6.settings/v1", "bidMode": bid_mode}, ensure_ascii=False),
            encoding="utf-8")

    def _material(self):
        return {"personnel": [{"name": "王监理"}],
                "performance": [{"name": "某市政改造监理项目"}]}

    def test_mingbiao_names_present_passes(self):
        self._settings("明标")
        _mk_docx(self.f, names="总监理工程师：王监理，负责某市政改造监理项目。")
        ok, issues = R.check_tech_ref_consistency(str(self.f), str(self.proj), self._material(), {})
        self.assertTrue(ok, issues)

    def test_mingbiao_missing_name_flagged(self):
        self._settings("明标")
        _mk_docx(self.f, names="正文未提任何业绩。")
        ok, issues = R.check_tech_ref_consistency(str(self.f), str(self.proj), self._material(), {})
        self.assertFalse(ok)
        joined = "\n".join(issues)
        self.assertIn("王监理", joined)
        self.assertIn("某市政改造监理项目", joined)

    def test_anbiao_skips(self):
        self._settings("暗标")
        _mk_docx(self.f)
        ok, issues = R.check_tech_ref_consistency(str(self.f), str(self.proj), self._material(), {})
        self.assertTrue(ok)
        self.assertTrue(any("跳过" in x for x in issues))

    def test_no_names_skips(self):
        self._settings("明标")
        _mk_docx(self.f)
        ok, issues = R.check_tech_ref_consistency(str(self.f), str(self.proj),
                                                  {"personnel": [], "performance": []}, {})
        self.assertTrue(ok)
        self.assertTrue(any("跳过" in x for x in issues))


class RemedyLinkTest(unittest.TestCase):
    def _check(self, ok_items=None, tech=True):
        ok_items = [{"id": "tech_structure", "name": "技术标结构完整性", "ok": True, "issues": []},
                    {"id": "tech_blank_pages", "name": "空白页检查", "ok": False,
                     "issues": ["第 3 段分页符所在段为空（疑空白页）"]},
                    {"id": "tech_figure_placement", "name": "图表位置合规", "ok": False,
                     "issues": ["正文中段插入图表（段落 9）"]}]
        return {"kind": "tech" if tech else "biz", "tag": "技术标" if tech else "",
                "检查": ok_items}

    def test_tasks_generated_for_failed_checks(self):
        tasks = remedy.gen_check_tasks(self._check())
        self.assertEqual(len(tasks), 2)
        self.assertTrue(all(t["类型"] == "格式合规" for t in tasks))
        self.assertTrue(all(t["动作"] == "技术标检查整改" for t in tasks))
        self.assertEqual(tasks[0]["来源"], "M7-技术标")
        self.assertIn("修复", tasks[0]["动作步骤"][0])
        self.assertIn("check-rule", tasks[0]["验证方式"])

    def test_passed_checks_skipped(self):
        tasks = remedy.gen_check_tasks({"kind": "tech", "tag": "技术标", "检查": [
            {"id": "x", "name": "通过项", "ok": True, "issues": []}]})
        self.assertEqual(tasks, [])

    def test_biz_kind_wording(self):
        tasks = remedy.gen_check_tasks(self._check(tech=False))
        self.assertEqual(tasks[0]["动作"], "标书检查整改")

    def test_empty_check_doc(self):
        self.assertEqual(remedy.gen_check_tasks({}), [])
        self.assertEqual(remedy.gen_check_tasks(None), [])


if __name__ == "__main__":
    unittest.main()
